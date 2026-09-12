"""Resume-safe, synthetic business journey; real requests are recorded before dispatch.

Never prints credentials or sends user documents. Existing recorded requests are
not repeated automatically, even if a previous run was interrupted.
"""
from argparse import ArgumentParser
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from document_pipeline_api.config import Settings
from document_pipeline_api.main import create_app
from document_pipeline_api.model_providers.openai_compatible import OpenAICompatibleProvider
from document_pipeline_api.services.extraction import process_task


def main():
    parser = ArgumentParser()
    parser.add_argument("--key-file", type=Path, required=True)
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[1] / ".local/acceptance-design-live"
    root.mkdir(parents=True, exist_ok=True)
    ledger = root / "journey.json"
    state = json.loads(ledger.read_text(encoding="utf-8")) if ledger.exists() else {"requests": [], "cases": {}}
    settings = Settings(database_url=f"sqlite:///{root / 'document-pipeline.db'}", storage_dir=root / "uploads", queue_enabled=False)
    stage = "setup"

    def record():
        temporary = ledger.with_suffix(".tmp")
        temporary.write_bytes(json.dumps(state, ensure_ascii=False, indent=2).encode())
        temporary.replace(ledger)

    original_post = OpenAICompatibleProvider._post_completion

    def counted_post(provider, payload, result_type):
        identity = f"{stage}:{result_type.__name__}"
        assert not any(item["identity"] == identity for item in state["requests"]), f"Recorded request {identity} will not be repeated; inspect its result first."
        prompt = json.dumps(payload.get("messages", []), ensure_ascii=False)
        if stage == "prefix":
            assert "TAIL_NOT_SENT_917" not in prompt
        item = {"identity": identity, "time": datetime.now(timezone.utc).isoformat(), "model": provider.model_name,
                "status": "dispatched", "prompt_sha256": sha256(prompt.encode()).hexdigest(), "characters": len(prompt)}
        state["requests"].append(item)
        record()
        try:
            result = original_post(provider, payload, result_type)
            item["status"] = "returned"
            return result
        except Exception as error:
            item["status"] = "failed"
            item["error_type"] = type(error).__name__
            raise
        finally:
            record()

    with TestClient(create_app(settings)) as client, patch.object(OpenAICompatibleProvider, "_post_completion", counted_post):
        def req(method, path, **kwargs):
            response = client.request(method, "/api/v1" + path, **kwargs)
            if not response.is_success:
                raise RuntimeError(f"{method} {path}: HTTP {response.status_code}: {response.text[:500]}")
            return response

        def read_task(task_id):
            return next(task for task in req("GET", "/tasks").json() if task["id"] == task_id)

        if "profile" not in state:
            profile = req("POST", "/models/profiles", json={"name": "合成全链路验收", "provider": "openai_compatible",
                "base_url": "https://dashscope.aliyuncs.com/compatible-mode/v1", "model_name": "qwen3.6-flash",
                "timeout_seconds": 90, "reasoning_effort": None, "api_key": args.key_file.read_text(encoding="utf-8-sig").strip(),
                "acknowledge_remote_data_transfer": True}).json()
            state["profile"] = {"id": profile["id"], "version": profile["version"]}
            record()
        profile = state["profile"]
        req("POST", f"/models/profiles/{profile['id']}/activate", json={"version": profile["version"], "acknowledge_remote_data_transfer": True})

        stage = "generate"
        if "draft" not in state:
            state["draft"] = req("POST", "/templates/generate", data={"requirement": "创建通用会议纪要模板。只要四个每份文件一次的文本字段，名称依次为：会议名称、日期、主要决定、后续行动。默认卡片阅读。不生成示例或校验规则。", "with_examples": "false", "with_rules": "false", "profile_id": profile["id"]}).json()
            record()
        if "template" not in state:
            draft = state["draft"]
            assert len(draft["fields"]) == 4
            created = req("POST", "/templates", json={"name": "通用会议纪要（真实 AI 验收）", "description": draft["description"],
                "fields": draft["fields"], "behavior": {**draft["behavior"], "suggest_filename": True}}).json()
            state["template"] = created
            record()
        template = state["template"]
        if not state.get("reordered"):
            fields = template["fields"]
            body = {key: template[key] for key in ["name", "description", "fields", "extra_instructions", "validation_rules", "deterministic_rules", "output_mapping", "behavior"]}
            body["fields"] = [fields[0], fields[2], fields[3], fields[1]]
            saved = req("PUT", f"/templates/{template['id']}", json={**body, "expected_version": template["version"]}).json()
            reloaded = req("GET", f"/templates/{template['id']}").json()
            assert reloaded["fields"] == body["fields"]
            state.update(template=saved, reordered=True)
            record()
            template = saved

        content = "合成测试材料，无真实用户信息。\n会议名称：研发评审会议\n日期：2026年9月12日\n主要决定：先验证用户流程，再发布升级版本。\n后续行动：整理试用反馈，核对原文件与结构化结果。\n"
        for limited, filename in [(False, "研发评审会议纪要-20260912.txt"), (True, "1001656.txt")]:
            stage = "prefix" if limited else "complete"
            if state["cases"].get(stage, {}).get("passed"):
                continue
            case = state["cases"].setdefault(stage, {})
            raw = (content + ("补充观察记录。" * 200 + "TAIL_NOT_SENT_917" if limited else "")).encode()
            (root / filename).write_bytes(raw)
            req("PUT", "/system/settings", json={"allow_limited_input": limited, "input_text_limit": 180})
            external = root / "external-copy"
            external.mkdir(exist_ok=True)
            if limited:
                binding_url = f"/templates/{template['id']}/local-export"
                binding = req("GET", binding_url).json()
                req("PUT", binding_url, json={"expected_revision": binding["revision"], "enabled": True, "parent_path": str(external)})
            if "task_id" not in case:
                form = {"template_id": template["id"]} if limited else {"template_mode": "smart"}
                task = req("POST", "/tasks", files={"file": (filename, raw, "text/plain")}, data=form).json()
                case["task_id"] = task["id"]
                record()
            task_id = case["task_id"]
            task = read_task(task_id)
            if task["status"] not in {"needs_review", "completed"}:
                with client.app.state.session_factory() as session:
                    extracted = process_task(session, settings, task_id)
                    assert extracted is not None, read_task(task_id).get("failure_message")
            result = req("GET", f"/tasks/{task_id}/result").json()
            case["result"] = result
            record()
            assert result["template_id"] == template["id"], "Smart match must select the generated template"
            assert result["input_scope"]["coverage"] == ("partial" if limited else "complete")
            suggestion = result["file_name"]["suggested_filename"]
            assert suggestion != filename if limited else suggestion == filename
            assert req("GET", f"/tasks/{task_id}/file").content == raw
            conflict = external / template["name"] / suggestion
            if limited and not case.get("confirmed"):
                conflict.parent.mkdir(parents=True, exist_ok=True)
                conflict.write_bytes(b"existing independent file - do not overwrite")
            if not case.get("confirmed"):
                confirmation = req("POST", f"/tasks/{task_id}/confirm", json={"expected_review_version": result["review_version"], "filename": suggestion}).json()
                case.update(confirmed=True, table_id=confirmation["table_id"])
                record()
            table = req("GET", f"/tables/{case['table_id']}").json()
            assert table["rows"]
            assert [col["key"] for col in table["columns"]] == [field["key"] for field in template["fields"]]
            workbook = load_workbook(BytesIO(req("GET", f"/tables/{case['table_id']}/export.xlsx").content))
            assert workbook.active.max_row >= 2
            case["export_headers"] = [cell.value for cell in workbook.active[1]]
            if limited:
                task = read_task(task_id)
                assert task["status"] == "completed"
                if not case.get("recovered"):
                    assert task["file_export"]["error_code"] == "name_conflict"
                    count = len(state["requests"])
                    recovery = root / "recovered-copy"
                    recovery.mkdir(exist_ok=True)
                    recovered = req("POST", f"/tasks/{task_id}/export", json={"action": "retry", "filename": suggestion, "parent_path": str(recovery)}).json()
                    assert recovered["file_export"]["status"] == "completed"
                    assert Path(recovered["file_export"]["actual_path"]).read_bytes() == raw
                    assert conflict.read_bytes() == b"existing independent file - do not overwrite"
                    assert len(state["requests"]) == count
                    case.update(recovered=True, copy=recovered["file_export"]["actual_path"])
            case.update(passed=True, original_sha256=sha256(raw).hexdigest())
            record()
        state["passed"] = True
        record()
    print(json.dumps({"passed": state.get("passed", False), "requests": len(state["requests"]), "ledger": str(ledger)}))


if __name__ == "__main__":
    main()
