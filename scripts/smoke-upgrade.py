"""Synthetic, isolated live-model acceptance. Never reads application user data.

Run with the API project's Python: scripts/smoke-upgrade.py --key-file <local file>.
Only generated fixtures are sent to the configured provider. Credentials stay in memory.
"""
from argparse import ArgumentParser
from dataclasses import replace
from datetime import datetime
import json
from pathlib import Path
import zipfile

from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont

from document_pipeline_api.config import Settings
from document_pipeline_api.main import create_app
from document_pipeline_api.model_providers import build_model_provider
from document_pipeline_api.services.extraction import process_task


def fixtures(directory: Path) -> list[tuple[Path, str, bool]]:
    directory.mkdir(parents=True)
    text = "文档名称：校园科学展览\n" + "\n".join(
        f"活动记录 {i}：同学们观察植物生长并记录实验数据。" for i in range(1000)
    ) + "\n归档编号：ZX-0908\n"
    (directory / "长记录.md").write_text(text, encoding="utf-8")
    (directory / "1001656.txt").write_text(
        "科目：数学\n第一题：计算 7 + 8。\n正确答案：15。\n错因：把加法看成减法。\n\n"
        "第二题：计算 6 × 9。\n正确答案：54。\n错因：乘法口诀记错。\n", encoding="utf-8")
    book = Workbook()
    first = book.active
    first.title = "活动说明"
    first.append(["文档名称", "校园科学展览"])
    sheet = book.create_sheet("实验记录")
    sheet.append(["序号", "观察"])
    for i in range(20000):
        sheet.append([i + 1, "植物生长记录"])
    book.create_sheet("归档信息").append(["归档编号", "ZX-0908"])
    book.save(directory / "多表两万行.xlsx")
    with zipfile.ZipFile(directory / "活动卡片.docx", "w") as archive:
        archive.writestr("word/document.xml", '''<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>文档名称：校园科学展览</w:t></w:r></w:p><w:tbl><w:tr><w:tc><w:p><w:r><w:t>归档编号</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>ZX-0908</w:t></w:r></w:p></w:tc></w:tr></w:tbl></w:body></w:document>''')
        archive.writestr("[Content_Types].xml", '''<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>''')
        archive.writestr("_rels/.rels", '''<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>''')
    font = ImageFont.truetype("C:/Windows/Fonts/msyh.ttc", 30)
    pages = []
    for i in range(60):
        page = Image.new("RGB", (700, 900), "white")
        draw = ImageDraw.Draw(page)
        draw.text((35, 35), f"展览记录 · 第 {i + 1} 页", font=font, fill="black")
        if i == 0:
            draw.text((35, 120), "文档名称：校园科学展览", font=font, fill="black")
        elif i == 59:
            draw.text((35, 120), "归档编号：ZX-0908", font=font, fill="black")
        else:
            draw.text((35, 120), "同学们观察植物并记录。", font=font, fill="black")
        pages.append(page)
    pages[0].save(directory / "六十页展览记录.pdf", save_all=True, append_images=pages[1:])
    for page in pages:
        page.close()
    return [
        (directory / "长记录.md", "text/markdown", True),
        (directory / "多表两万行.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", True),
        (directory / "活动卡片.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", False),
        (directory / "六十页展览记录.pdf", "application/pdf", True),
    ]


def main():
    parser = ArgumentParser(description=__doc__)
    parser.add_argument("--key-file", type=Path, required=True)
    parser.add_argument("--base-url", default="https://dashscope.aliyuncs.com/compatible-mode/v1")
    parser.add_argument("--model", default="qwen3.6-flash")
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[1]
    run = root / ".local/tests" / ("upgrade-live-" + datetime.now().strftime("%Y%m%d-%H%M%S-%f"))
    cases = fixtures(run / "测试文件")
    settings = Settings(database_url=f"sqlite:///{run / 'acceptance.db'}", storage_dir=run / "uploads",
        model_provider="openai_compatible", model_name=args.model, model_base_url=args.base_url,
        model_reasoning_effort=None, model_timeout_seconds=60, max_pdf_pages=4, queue_enabled=False)
    key = args.key_file.read_text(encoding="utf-8-sig").strip()
    provider = build_model_provider(replace(settings, model_api_key=key), timeout_seconds=60)
    report = {"model": args.model, "cases": [], "success": False}
    try:
        with TestClient(create_app(settings)) as client:
            def request(method, path, **kwargs):
                response = client.request(method, "/api/v1" + path, **kwargs)
                response.raise_for_status()
                return response

            template = request("POST", "/templates", json={"name": "活动元信息验收",
                "description": "提取活动文档元信息，不要求穷尽活动记录。",
                "fields": [{"key": "title", "label": "文档名称", "section": "header"},
                           {"key": "archive_id", "label": "归档编号", "section": "header"}],
                "behavior": {"requires_complete_input": False, "presentation": {"mode": "card", "title_field": "header.title"}}}).json()
            external = run / "外部副本"
            external.mkdir()
            request("PUT", f"/templates/{template['id']}/local-export", json={
                "expected_revision": 0, "enabled": True, "parent_path": str(external)})
            first_path, first_type, _ = cases[0]
            rejected = client.post("/api/v1/tasks", files={"file": (first_path.name, first_path.read_bytes(), first_type)}, data={"template_id": template["id"]})
            assert rejected.status_code == 422, "Default must reject oversized model input"
            assert request("GET", "/tasks").json() == []
            report["default_rejection"] = True
            request("PUT", "/system/settings", json={"image_convert": True, "office_convert": True, "allow_limited_input": True})
            for index, (path, mime, partial) in enumerate(cases):
                evidence = {"file": str(path), "success": False}
                report["cases"].append(evidence)
                try:
                    task = request("POST", "/tasks", files={"file": (path.name, path.read_bytes(), mime)}, data={"template_id": template["id"]}).json()
                    with client.app.state.session_factory() as session:
                        result = process_task(session, settings, task["id"], client=provider)
                    assert result is not None, "Extraction missing"
                    assert result.result.header["archive_id"] == "ZX-0908", "Tail metadata not extracted"
                    assert "校园科学展览" in result.result.header["title"], "Head metadata not extracted"
                    assert result.input_scope.coverage == ("partial" if partial else "complete")
                    assert request("GET", f"/tasks/{task['id']}/file").content == path.read_bytes()
                    destination = external / template["name"] / path.name
                    if index == 0:
                        destination.parent.mkdir(parents=True)
                        destination.write_bytes(b"unrelated existing user file")
                    confirmation = request("POST", f"/tasks/{task['id']}/confirm", json={"expected_review_version": 0}).json()
                    table = request("GET", f"/tables/{confirmation['table_id']}").json()
                    assert any(row["task_id"] == task["id"] and row["input_scope"]["coverage"] == result.input_scope.coverage for row in table["rows"])
                    request("GET", f"/tables/{confirmation['table_id']}/export.xlsx")
                    if index == 0:
                        pending = request("GET", "/tasks?export_pending=true").json()
                        assert any(row["id"] == task["id"] and row["status"] == "completed" for row in pending)
                        alternative = run / "改选副本目录"
                        alternative.mkdir()
                        state = request("POST", f"/tasks/{task['id']}/export", json={"action": "retry", "parent_path": str(alternative), "filename": "展览副本.md"}).json()
                        assert state["file_export"]["status"] == "completed"
                        assert destination.read_bytes() == b"unrelated existing user file"
                        assert Path(state["file_export"]["actual_path"]).read_bytes() == path.read_bytes()
                    else:
                        assert destination.read_bytes() == path.read_bytes()
                    evidence.update(success=True, task_id=task["id"], table_id=confirmation["table_id"], result=result.result.model_dump(), scope=result.input_scope.model_dump(), seconds=result.elapsed_seconds)
                except Exception as error:
                    evidence.update(error_type=type(error).__name__, error=str(error).replace(key, "[REDACTED]")[:2000])
                print(json.dumps({"file": path.name, "success": evidence["success"], "error": evidence.get("error")}, ensure_ascii=False), flush=True)
                (run / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
            report["success"] = all(case["success"] for case in report["cases"])
    finally:
        provider.close()
        (run / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps({"report": str(run / "report.json"), "success": report["success"]}, ensure_ascii=False))
    return 0 if report["success"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
