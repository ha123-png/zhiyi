"""Bounded local integrity/load checks; no model, credentials or user documents."""

from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import time
import uuid
import zipfile

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from PIL import Image
import pypdfium2 as pdfium

from document_pipeline_api.config import Settings
from document_pipeline_api.main import create_app
from document_pipeline_api.services.extraction import process_task


class Model:
    model_name = "offline-pressure-fixture"
    calls = 0

    def complete_text(self, prompt, result_type):
        self.calls += 1
        assert "TAIL_NOT_SENT" not in prompt
        assert len(prompt) < 10000
        return result_type.model_validate(
            {"header": {"title": "合成材料范围核对"}, "items": []}
        )

    def extract_images(self, images, prompt, result_type):
        assert 1 <= len(images) <= 4
        assert all(path.is_file() for path in images)
        return self.complete_text(prompt, result_type)

    def extract_image(self, image, prompt, result_type):
        return self.extract_images([image], prompt, result_type)


def main():
    root = (
        Path(__file__).resolve().parents[1]
        / ".local/release-pressure"
        / uuid.uuid4().hex[:12]
    )
    root.mkdir(parents=True)
    (root.parent / "latest.json").write_bytes(json.dumps({"root": str(root)}).encode())
    report = {"passed": False, "real_model_calls": 0, "cases": [], "queue": {}}

    def save():
        (root / "report.json").write_bytes(
            json.dumps(report, ensure_ascii=False, indent=2).encode()
        )

    settings = Settings(
        database_url=f"sqlite:///{root / 'document-pipeline.db'}",
        storage_dir=root / "uploads",
        max_active_tasks=40,
    )
    model = Model()
    files = []
    # Long first line plus many lines tests character and row caps together.
    text = (
        "开头"
        + "甲" * 10000
        + "\n"
        + ("合成正文，逐行验证完整原件。\n" * 90000)
        + "TAIL_NOT_SENT"
    ).encode()
    files.append(("large.txt", "text/plain", text, "90,002 lines; long first line"))
    files.append(("large.md", "text/markdown", text, "same long text as Markdown"))
    output = BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        xml = '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
        xml += "<w:p><w:r><w:t>" + "长首段。" * 3000 + "</w:t></w:r></w:p>"
        xml += "<w:p><w:r><w:t>合成正文段落。</w:t></w:r></w:p>" * 20000
        xml += "<w:p><w:r><w:t>TAIL_NOT_SENT</w:t></w:r></w:p></w:body></w:document>"
        archive.writestr("word/document.xml", xml)
    files.append(
        (
            "large.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            output.getvalue(),
            "20,002 paragraphs; long first paragraph",
        )
    )
    workbook = Workbook(write_only=True)
    for number in range(3):
        sheet = workbook.create_sheet(f"分组{number + 1}")
        sheet.append(["编号", "项目", "数量"])
        for row in range(12000):
            sheet.append([row, "合成物资", row % 7])
        sheet.append(["TAIL_NOT_SENT"])
    output = BytesIO()
    workbook.save(output)
    files.append(
        (
            "large.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            output.getvalue(),
            "3 sheets; 36,006 total rows",
        )
    )
    pdf = pdfium.PdfDocument.new()
    for _ in range(80):
        page = pdf.new_page(595, 842)
        page.close()
    output = BytesIO()
    pdf.save(output)
    pdf.close()
    files.append(
        ("large.pdf", "application/pdf", output.getvalue(), "80 pages; 4 rendered")
    )
    frame = Image.new("RGB", (1500, 1000), "white")
    output = BytesIO()
    frame.save(
        output,
        format="TIFF",
        save_all=True,
        append_images=[frame] * 19,
        compression="tiff_lzw",
    )
    frame.close()
    files.append(
        ("large.tiff", "image/tiff", output.getvalue(), "20 frames; 30 million pixels")
    )
    with TestClient(create_app(settings)) as client:

        def req(method, path, **kwargs):
            response = client.request(method, "/api/v1" + path, **kwargs)
            assert response.is_success, (
                f"{method} {path}: {response.status_code}: {response.text[:200]}"
            )
            return response

        template = req(
            "POST",
            "/templates",
            json={
                "name": "压力范围概览",
                "fields": [{"key": "title", "label": "标题"}],
            },
        ).json()
        req(
            "PUT",
            "/system/settings",
            json={
                "office_convert": True,
                "image_convert": True,
                "allow_limited_input": True,
                "input_text_limit": 2048,
                "input_page_limit": 4,
                "input_row_limit": 100,
            },
        )
        for filename, mime, raw, scale in files:
            started = time.monotonic()
            case = {
                "file": filename,
                "bytes": len(raw),
                "scale": scale,
                "source_sha256": sha256(raw).hexdigest(),
            }
            report["cases"].append(case)
            save()
            source = root / filename
            source.write_bytes(raw)
            task = req(
                "POST",
                "/tasks",
                files={"file": (filename, raw, mime)},
                data={"template_id": template["id"]},
            ).json()
            case["task_id"] = task["id"]
            save()
            assert task["planned_scope"]["coverage"] == "partial"
            with client.app.state.session_factory() as session:
                result = process_task(session, settings, task["id"], client=model)
                assert result is not None
            assert result.input_scope.coverage == "partial"
            assert any(
                issue.code == "partial_input" for issue in result.validation_issues
            )
            assert (
                sha256(req("GET", f"/tasks/{task['id']}/file").content).hexdigest()
                == case["source_sha256"]
            )
            confirmation = req(
                "POST",
                f"/tasks/{task['id']}/confirm",
                json={"expected_review_version": result.review_version},
            ).json()
            # Confirmation and repeat confirmation must not duplicate a data row.
            repeated = req(
                "POST",
                f"/tasks/{task['id']}/confirm",
                json={"expected_review_version": result.review_version},
            ).json()
            assert repeated["table_id"] == confirmation["table_id"]
            case.update(
                passed=True,
                elapsed_seconds=round(time.monotonic() - started, 3),
                scope=result.input_scope.model_dump(mode="json"),
            )
            save()
        table_id = confirmation["table_id"]
        table = req("GET", f"/tables/{table_id}?page_size=100").json()
        assert table["row_count"] == len(files)
        exported = req("GET", f"/tables/{table_id}/export.xlsx").content
        workbook = load_workbook(BytesIO(exported))
        assert workbook.active.max_row == len(files) + 1
        started = time.monotonic()
        ids = []
        for number in range(40):
            task = req(
                "POST",
                "/tasks",
                files={
                    "file": (
                        f"batch-{number}.txt",
                        f"合成批处理文件 {number}".encode(),
                        "text/plain",
                    )
                },
                data={"template_id": template["id"]},
            ).json()
            ids.append(task["id"])
        rejected = client.post(
            "/api/v1/tasks",
            files={"file": ("over-cap.txt", b"bounded queue", "text/plain")},
            data={"template_id": template["id"]},
        )
        assert rejected.status_code in {409, 429}, rejected.status_code
        report["queue"] = {
            "admitted": len(ids),
            "over_capacity_status": rejected.status_code,
        }
        for task_id in ids:
            with client.app.state.session_factory() as session:
                assert (
                    process_task(session, settings, task_id, client=model) is not None
                )
        report["queue"].update(
            processed=40, elapsed_seconds=round(time.monotonic() - started, 3)
        )
        report["offline_fixture_calls"] = model.calls
        report["passed"] = True
        save()
    print(
        json.dumps(
            {
                "passed": True,
                "formats": len(files),
                "queue": report["queue"],
                "root": str(root),
            }
        )
    )


if __name__ == "__main__":
    main()
