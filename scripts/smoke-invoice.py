"""Run one real-model invoice regression using only a generated synthetic image."""
from argparse import ArgumentParser
from dataclasses import replace
from datetime import datetime
from io import BytesIO
import json
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont

from document_pipeline_api.config import Settings
from document_pipeline_api.main import create_app
from document_pipeline_api.model_providers import build_model_provider
from document_pipeline_api.services.extraction import process_task


def main():
    parser = ArgumentParser(description=__doc__)
    parser.add_argument("--key-file", type=Path, required=True)
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[1]
    run = root / ".local/tests" / ("invoice-live-" + datetime.now().strftime("%Y%m%d-%H%M%S-%f"))
    run.mkdir(parents=True)
    path = run / "合成发票验收.png"
    font = ImageFont.truetype("C:/Windows/Fonts/msyh.ttc", 32)
    with Image.new("RGB", (1300, 1000), "white") as image:
        draw = ImageDraw.Draw(image)
        for index, line in enumerate([
            "发票（合成测试样例，无真实交易，不可报销）",
            "发票号码：TEST09080001    开票日期：2026-09-08",
            "销售方：星光文具测试公司",
            "购买方：春苗学校测试单位",
            "商品名称    规格    单位    数量    不含税单价    金额    税率    税额",
            "练习本        A5      本       10            5.00          50.00     6%      3.00",
            "文件夹        A4      个        5           10.00         50.00     6%      3.00",
            "不含税金额合计：100.00 元",
            "税额合计：6.00 元",
            "价税合计：106.00 元",
        ]):
            draw.text((40, 40 + index * 80), line, font=font, fill="black")
        image.save(path)
    settings = Settings(database_url=f"sqlite:///{run / 'acceptance.db'}", storage_dir=run / "uploads",
        model_provider="openai_compatible", model_name="qwen3.6-flash",
        model_base_url="https://dashscope.aliyuncs.com/compatible-mode/v1",
        model_reasoning_effort=None, model_timeout_seconds=60, queue_enabled=False)
    key = args.key_file.read_text(encoding="utf-8-sig").strip()
    provider = build_model_provider(replace(settings, model_api_key=key), timeout_seconds=60)
    report = {"success": False, "file": str(path), "model": settings.model_name}
    try:
        with TestClient(create_app(settings)) as client:
            def request(method, endpoint, **kwargs):
                response = client.request(method, "/api/v1" + endpoint, **kwargs)
                response.raise_for_status()
                return response

            task = request("POST", "/tasks", files={"file": (path.name, path.read_bytes(), "image/png")},
                data={"template_id": "builtin-invoice"}).json()
            with client.app.state.session_factory() as session:
                result = process_task(session, settings, task["id"], client=provider)
            assert result is not None
            report["extraction"] = result.model_dump(mode="json")
            assert result.result.document_number == "TEST09080001"
            assert result.result.total_amount == 106
            assert result.result.amount_before_tax == 100 and result.result.tax_amount == 6
            assert len(result.result.items) == 2
            assert sorted(item.amount for item in result.result.items) == [50, 50]
            assert result.validation_issues == []
            assert result.input_scope.coverage == "complete"
            assert any(item.page_number == 1 and item.location_verified for item in result.evidence)
            confirmation = request("POST", f"/tasks/{task['id']}/confirm", json={"expected_review_version": 0}).json()
            table = request("GET", f"/tables/{confirmation['table_id']}").json()
            assert table["row_count"] == 2 and table["presentation"]["mode"] == "table"
            assert all(row["task_id"] == task["id"] for row in table["rows"])
            exported = request("GET", f"/tables/{confirmation['table_id']}/export.xlsx").content
            (run / "合成发票结果.xlsx").write_bytes(exported)
            workbook = load_workbook(BytesIO(exported), read_only=True, data_only=True)
            try:
                values = [value for sheet in workbook for row in sheet.iter_rows(values_only=True) for value in row]
                assert "TEST09080001" in values and 106 in values
            finally:
                workbook.close()
            assert request("GET", f"/tasks/{task['id']}/file").content == path.read_bytes()
            report.update(success=True, task_id=task["id"], table_id=table["id"], rows=table["rows"])
    except Exception as error:
        report.update(error_type=type(error).__name__, error=str(error).replace(key, "[REDACTED]")[:2000])
    finally:
        provider.close()
        (run / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"success": report["success"], "report": str(run / "report.json"), "error": report.get("error")}, ensure_ascii=True))
    return 0 if report["success"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
