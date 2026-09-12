"""Synthetic real-provider checks with a durable, replay-safe request ledger.

Uses only the selected existing model and an explicitly supplied key file. Outputs
are synthetic; credentials and HTTP headers are never serialized.
"""

import ast
from argparse import ArgumentParser
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import subprocess
import time
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont

from document_pipeline_api.config import Settings
from document_pipeline_api.main import create_app
from document_pipeline_api.model_providers.openai_compatible import (
    OpenAICompatibleProvider,
)
from document_pipeline_api.schemas.templates import TemplateRead
from document_pipeline_api.services.extraction import process_task
from document_pipeline_api.services.template_ai import generate_template_draft
from document_pipeline_api.services.template_processing import (
    build_template_extraction_model,
)


def main():
    parser = ArgumentParser()
    parser.add_argument("--provider", choices=["lmstudio", "cloud"], required=True)
    parser.add_argument("--key-file", type=Path)
    parser.add_argument("--generation-revision", default="original")
    parser.add_argument("--with-image", action="store_true")
    args = parser.parse_args()
    root = (
        Path(__file__).resolve().parents[1]
        / ".local/release-model-check"
        / args.provider
    )
    root.mkdir(parents=True, exist_ok=True)
    ledger = root / "ledger.json"
    state = (
        json.loads(ledger.read_text(encoding="utf-8"))
        if ledger.exists()
        else {"calls": {}, "provider": args.provider}
    )

    def save():
        temp = ledger.with_suffix(".tmp")
        temp.write_bytes(json.dumps(state, ensure_ascii=False, indent=2).encode())
        temp.replace(ledger)

    state["passed"] = False
    save()
    cloud = args.provider == "cloud"
    key = (
        args.key_file.read_text(encoding="utf-8-sig").strip()
        if cloud and args.key_file
        else ""
    )
    if cloud and not key:
        raise ValueError("Cloud validation requires the explicitly supplied key file")
    provider = OpenAICompatibleProvider(
        "https://dashscope.aliyuncs.com/compatible-mode/v1"
        if cloud
        else "http://127.0.0.1:1234/v1",
        "qwen3.6-flash" if cloud else "zhiyi-acceptance-qwen4b",
        api_key=key,
        temperature=0,
        timeout_seconds=120,
    )
    settings = Settings(
        database_url=f"sqlite:///{root / 'document-pipeline.db'}",
        storage_dir=root / "uploads",
    )
    stage = "setup"
    original = OpenAICompatibleProvider._post_completion

    def counted(instance, payload, result_type):
        identity = stage + ":" + result_type.__name__
        encoded = json.dumps(payload, ensure_ascii=False).encode()
        digest = sha256(encoded).hexdigest()
        prior = state["calls"].get(identity)
        if prior:
            if prior.get("status") == "returned":
                assert prior["request_sha256"] == digest, (
                    "Request changed; cached evidence must not be silently reused"
                )
                return result_type.model_validate(prior["result"])
            raise RuntimeError(
                "This request was already dispatched; inspect its evidence before retrying"
            )
        entry = {
            "status": "dispatched",
            "request_sha256": digest,
            "request_bytes": len(encoded),
            "model": instance.model_name,
            "at": datetime.now(timezone.utc).isoformat(),
        }
        state["calls"][identity] = entry
        save()
        started = time.monotonic()
        try:
            result = original(instance, payload, result_type)
            entry.update(status="returned", result=result.model_dump(mode="json"))
            return result
        except Exception as error:
            entry.update(status="failed", error_type=type(error).__name__)
            raise
        finally:
            entry["elapsed_seconds"] = round(time.monotonic() - started, 3)
            save()

    with (
        TestClient(create_app(settings)) as client,
        patch.object(OpenAICompatibleProvider, "_post_completion", counted),
    ):

        def req(method, path, **kwargs):
            response = client.request(method, "/api/v1" + path, **kwargs)
            assert response.is_success, f"{method} {path}: {response.status_code}"
            return response

        if "template" not in state:
            state["template"] = req(
                "POST",
                "/templates",
                json={
                    "name": "物资领用核对（合成验证）",
                    "fields": [
                        {"key": "title", "label": "领用单名称", "section": "header"},
                        {"key": "owner", "label": "审批人", "section": "header"},
                        {"key": "name", "label": "物品", "section": "item"},
                        {
                            "key": "qty",
                            "label": "数量",
                            "section": "item",
                            "value_type": "number",
                        },
                        {
                            "key": "price",
                            "label": "单价",
                            "section": "item",
                            "value_type": "number",
                        },
                        {
                            "key": "amount",
                            "label": "金额",
                            "section": "item",
                            "value_type": "number",
                        },
                    ],
                },
            ).json()
            save()
        template = TemplateRead.model_validate(state["template"])
        raw = "合成测试文件，无真实用户数据。\n领用单名称：读书会活动物资\n明细：\n物品 | 数量 | 单价 | 金额\n便签 | 2 | 3.5 | 7\n便签 | 1 | 3.5 | 3.5\n黑色笔 | 0 | 2 | 0\n审批人未填写。\n".encode()
        expected = {
            "header": {"title": "读书会活动物资", "owner": None},
            "items": [
                {"name": "便签", "qty": 2, "price": 3.5, "amount": 7},
                {"name": "便签", "qty": 1, "price": 3.5, "amount": 3.5},
                {"name": "黑色笔", "qty": 0, "price": 2, "amount": 0},
            ],
        }
        stage = "baseline"
        previous = subprocess.check_output(
            [
                "git",
                "show",
                "a5ee460:apps/api/src/document_pipeline_api/services/template_processing.py",
            ]
        ).decode()
        function = next(
            node
            for node in ast.parse(previous).body
            if isinstance(node, ast.FunctionDef)
            and node.name == "build_template_extraction_prompt"
        )
        namespace = {"json": json, "TemplateRead": TemplateRead}
        exec(
            compile(
                ast.Module(body=[function], type_ignores=[]), "baseline-prompt", "exec"
            ),
            namespace,
        )
        baseline = provider.complete_text(
            namespace["build_template_extraction_prompt"](template)
            + "\n"
            + raw.decode(),
            build_template_extraction_model(template),
        )
        state["baseline_exact"] = baseline.model_dump() == expected
        save()
        stage = "extraction"
        if "task" not in state:
            state["task"] = req(
                "POST",
                "/tasks",
                files={"file": ("读书会物资.txt", raw, "text/plain")},
                data={"template_id": template.id},
            ).json()["id"]
            save()
        task_id = state["task"]
        if "extraction" not in state:
            with client.app.state.session_factory() as session:
                extracted = process_task(session, settings, task_id, client=provider)
                assert extracted is not None, "Inspect the synthetic task diagnostic"
            state["extraction"] = req("GET", f"/tasks/{task_id}/result").json()
            save()
        result = state["extraction"]
        state["extraction_exact"] = result["result"] == expected
        assert state["extraction_exact"], (
            "Synthetic facts differ; preserve evidence and inspect before proceeding"
        )
        assert req("GET", f"/tasks/{task_id}/file").content == raw
        assert result["input_scope"]["coverage"] == "complete"
        if "confirmation" not in state:
            state["confirmation"] = req(
                "POST",
                f"/tasks/{task_id}/confirm",
                json={"expected_review_version": result["review_version"]},
            ).json()
            save()
        table_id = state["confirmation"]["table_id"]
        table = req("GET", f"/tables/{table_id}").json()
        assert len(table["rows"]) == 3
        export = req("GET", f"/tables/{table_id}/export.xlsx").content
        workbook = load_workbook(BytesIO(export))
        assert workbook.active.max_row == 4
        (root / "synthetic-export.xlsx").write_bytes(export)
        state["original_sha256"] = sha256(raw).hexdigest()
        stage = "rule_generation" + (
            ""
            if args.generation_revision == "original"
            else ":" + args.generation_revision
        )
        with client.app.state.session_factory() as session:
            draft = generate_template_draft(
                settings,
                session,
                requirement="为物资领用单创建模板，只要四个明细字段：物品（文本）、数量（数字）、单价（数字）、金额（数字）。用户明确规定：数量可以是零；金额应等于数量乘单价。只建议这一条计算校验作为提醒，不要求任何必填、范围或枚举规则。默认表格。",
                image_paths=[],
                with_examples=False,
                with_rules=True,
                model_client=provider,
            )
        state["draft"] = draft.model_dump(mode="json")
        rules = draft.rule_suggestions
        state["rules_contract"] = (
            len(draft.fields) == 4
            and len(rules) == 1
            and rules[0].status == "accepted"
            and rules[0].rule.kind == "equation"
        )
        state["rule_basis_present"] = bool(
            rules and "AI 建议依据" in rules[0].explanation
        )
        save()
        assert state["rules_contract"] and state["rule_basis_present"], (
            "Inspect rule suggestions; no extra model request is issued automatically"
        )
        if args.with_image:
            image_path = root / "synthetic-materials.png"
            if not image_path.exists():
                image = Image.new("RGB", (1050, 480), "white")
                draw = ImageDraw.Draw(image)
                font = ImageFont.truetype("C:/Windows/Fonts/msyh.ttc", 27)
                for number, line in enumerate(
                    [
                        "合成演示材料 · 无真实用户信息",
                        "领用单名称：读书会活动物资",
                        "审批人：未填写",
                        "物品              数量       单价       金额",
                        "便签              2          3.5        7",
                        "便签              1          3.5        3.5",
                        "黑色笔            0          2          0",
                    ]
                ):
                    draw.text((35, 25 + number * 57), line, font=font, fill="black")
                image.save(image_path)
                image.close()
            stage = "visual_extraction"
            image_raw = image_path.read_bytes()
            if "image_task" not in state:
                state["image_task"] = req(
                    "POST",
                    "/tasks",
                    files={"file": ("读书会物资.png", image_raw, "image/png")},
                    data={"template_id": template.id},
                ).json()["id"]
                save()
            if "image_extraction" not in state:
                with client.app.state.session_factory() as session:
                    visual = process_task(
                        session, settings, state["image_task"], client=provider
                    )
                    assert visual is not None, (
                        "Inspect synthetic visual task diagnostic"
                    )
                state["image_extraction"] = visual.model_dump(mode="json")
                save()
            state["image_exact"] = state["image_extraction"]["result"] == expected
            assert state["image_exact"]
            assert req("GET", f"/tasks/{state['image_task']}/file").content == image_raw
            state["image_sha256"] = sha256(image_raw).hexdigest()
            save()
        state["passed"] = True
        save()
    provider.close()
    print(
        json.dumps(
            {
                "provider": args.provider,
                "passed": state.get("passed", False),
                "calls": len(state["calls"]),
            }
        )
    )


if __name__ == "__main__":
    main()
