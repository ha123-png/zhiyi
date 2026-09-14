"""Isolated file-to-data-to-backup journey with explicit offline model fixtures.

Run with apps/api/.venv's Python. No keys, user data, real model or packaging.
Each invocation keeps a fresh .local directory and a durable JSON report.
"""

import csv
from datetime import date
from hashlib import sha256
from io import BytesIO, StringIO
import json
from pathlib import Path
import sqlite3
from uuid import uuid4

from fastapi.testclient import TestClient
from fastapi import HTTPException
from openpyxl import load_workbook

from document_pipeline_api.business_backup import (
    create_business_backup, inspect_business_backup, restore_business_backup,
)
from document_pipeline_api.config import Settings
from document_pipeline_api.main import create_app
from document_pipeline_api.model_providers.base import ModelServiceError
from document_pipeline_api.services.extraction import process_task


class OfflineExtraction:
    model_name = "offline-journey-fixture"
    calls = 0
    fail_next = False

    def complete_text(self, prompt, result_type):
        self.calls += 1
        assert "TAIL_MUST_STAY_LOCAL" not in prompt
        if self.fail_next:
            self.fail_next = False
            raise ModelServiceError("Synthetic transient failure; no network request.")
        return result_type.model_validate({
            "header": {"vendor": "合成供货方", "total": 300, "date": date.today().isoformat()},
            "items": [{"item": "合成物品甲", "amount": 100}, {"item": "合成物品乙", "amount": 200}],
        })


class OfflineConversation:
    plan = []
    calls = 0

    def __init__(self, settings, cancel):
        self.step = 0

    def close(self):
        pass

    def stream(self, messages, tools):
        type(self).calls += 1
        if self.step < len(self.plan):
            event = self.plan[self.step]
            self.step += 1
            yield event
        else:
            yield {"type": "text", "text": "合成流程已返回结果；这条回答来自离线测试替身。"}


def main():
    root = Path(__file__).resolve().parents[1] / ".local/offline-journey" / uuid4().hex[:12]
    root.mkdir(parents=True)
    (root.parent / "latest.json").write_text(json.dumps({"root": str(root)}), encoding="utf-8")
    report = {"passed": False, "real_model_calls": 0, "stages": []}

    def save(stage, **values):
        report["stages"].append({"stage": stage, **values})
        (root / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    def settings_for(directory):
        return Settings(database_url=f"sqlite:///{directory / 'document-pipeline.db'}",
                        storage_dir=directory / "uploads", queue_enabled=False)

    source = root / "source"
    settings = settings_for(source)
    originals = {}
    model = OfflineExtraction()
    save("started", root=str(root))
    with TestClient(create_app(settings)) as client:
        client.app.state.conversation_factory = OfflineConversation

        def request(method, path, **kwargs):
            response = client.request(method, "/api/v1" + path, **kwargs)
            assert response.is_success, f"{method} {path}: {response.status_code} {response.text[:1000]}"
            return response

        template = request("POST", "/templates", json={"name": "离线完整链验收", "fields": [
            {"key": "vendor", "label": "供应方", "section": "header"},
            {"key": "total", "label": "单据总额", "section": "header", "value_type": "number"},
            {"key": "date", "label": "日期", "section": "header", "value_type": "date"},
            {"key": "item", "label": "物品", "section": "item"},
            {"key": "amount", "label": "明细额", "section": "item", "value_type": "number"},
        ]}).json()
        external = root / "external"
        external.mkdir()
        binding_url = f"/templates/{template['id']}/local-export"
        request("PUT", binding_url, json={"expected_revision": 0, "enabled": True, "parent_path": str(external)})
        request("PUT", "/system/settings", json={"allow_limited_input": True, "input_text_limit": 1024})
        tasks = []
        for index, kind in enumerate(("cancel_then_reimport", "failure_retry", "partial")):
            raw = (f"离线合成文件 {index}：两条明细 100 和 200，总额 300。\n"
                   + ("有界材料正文。" * 500 + "TAIL_MUST_STAY_LOCAL" if kind == "partial" else "")).encode()
            task = request("POST", "/tasks", files={"file": (f"synthetic-{kind}.txt", raw, "text/plain")},
                           data={"template_id": template["id"]}).json()
            originals[task["id"]] = raw
            before_calls = model.calls
            if kind == "cancel_then_reimport":
                assert request("POST", f"/tasks/{task['id']}/cancel").json()["status"] == "cancelled"
                with client.app.state.session_factory() as session:
                    try:
                        process_task(session, settings, task["id"], client=model)
                    except HTTPException as error:
                        assert error.status_code == 409
                    else:
                        raise AssertionError("A cancelled task must refuse processing.")
                assert model.calls == before_calls
                assert request("GET", f"/tasks/{task['id']}/file").content == raw
                # Cancellation is terminal in the product; retry is for failures.
                assert client.post(f"/api/v1/tasks/{task['id']}/retry").status_code == 409
                task = request("POST", "/tasks", files={"file": (task["filename"], raw, "text/plain")},
                               data={"template_id": template["id"]}).json()
                originals[task["id"]] = raw
            tasks.append(task["id"])
            if kind == "failure_retry":
                model.fail_next = True
                with client.app.state.session_factory() as session:
                    try:
                        process_task(session, settings, task["id"], client=model)
                    except ModelServiceError:
                        pass
                    else:
                        raise AssertionError("The injected model failure must propagate.")
                current = next(item for item in request("GET", "/tasks").json() if item["id"] == task["id"])
                assert current["status"] == "failed"
                assert request("POST", f"/tasks/{task['id']}/retry").json()["status"] == "queued"
            with client.app.state.session_factory() as session:
                result = process_task(session, settings, task["id"], client=model)
            assert result is not None
            assert result.input_scope.coverage == ("partial" if kind == "partial" else "complete")
            assert request("GET", f"/tasks/{task['id']}/file").content == raw
            conflict = external / template["name"] / task["filename"]
            if index == 0:
                conflict.parent.mkdir(parents=True, exist_ok=True)
                conflict.write_bytes(b"independent existing file")
            confirmation = request("POST", f"/tasks/{task['id']}/confirm",
                                   json={"expected_review_version": result.review_version}).json()
            assert request("POST", f"/tasks/{task['id']}/confirm",
                           json={"expected_review_version": result.review_version}).json()["table_id"] == confirmation["table_id"]
            if index == 0:
                calls = model.calls
                alternative = root / "alternative"
                alternative.mkdir()
                state = request("POST", f"/tasks/{task['id']}/export", json={
                    "action": "retry", "parent_path": str(alternative), "filename": "synthetic-copy.txt",
                }).json()["file_export"]
                assert state["status"] == "completed" and Path(state["actual_path"]).read_bytes() == raw
                assert conflict.read_bytes() == b"independent existing file" and model.calls == calls
            save(kind, passed=True, task_id=task["id"], sha256=sha256(raw).hexdigest(), coverage=result.input_scope.coverage)

        table_id = confirmation["table_id"]
        table_url = f"/tables/{table_id}"
        table = request("GET", table_url).json()
        assert table["row_count"] == 6
        first = next(row for row in table["rows"] if row["task_id"] == tasks[0])
        updated = request("PATCH", f"{table_url}/rows/{first['id']}", json={
            "expected_version": first["version"], "changes": {"total": 350, "amount": 150},
        }).json()
        assert updated["version"] == first["version"] + 1
        conflict = client.patch("/api/v1" + f"{table_url}/rows/{first['id']}", json={
            "expected_version": first["version"], "changes": {"total": 999},
        })
        assert conflict.status_code == 409
        profile = request("GET", "/models/profiles").json()[0]

        def chat(tool, arguments, text):
            OfflineConversation.plan = [{"type": "tool", "id": uuid4().hex, "name": tool, "arguments": arguments}]
            started = request("POST", "/assistant/runs", json={"text": text, "profile_id": profile["id"],
                "profile_version": profile["version"], "context": {"table_id": table_id}}).json()
            assert "run.completed" in request("GET", f"/assistant/runs/{started['run_id']}/events").text
            thread = request("GET", f"/assistant/threads/{started['thread_id']}").json()
            assert thread["runs"][0]["status"] == "completed", thread
            return thread

        proposal = dict(table_id=table_id, row_id=first["id"], changes={"total": 400})
        rejected = chat("propose_changes", proposal, "把这张单据总额改为400，先让我核对。")
        request("POST", f"/assistant/tools/{rejected['tools'][0]['id']}/decision", json={"approve": False})
        approved = chat("propose_changes", proposal, "把这张单据总额改为400，先让我核对。")
        approval_url = f"/assistant/tools/{approved['tools'][0]['id']}/decision"
        assert request("POST", approval_url, json={"approve": True}).json()["status"] == "approved"
        assert request("POST", approval_url, json={"approve": True}).json()["status"] == "approved"
        analysis = chat("analyze_data_table", {"table_id": table_id, "metrics": [{"op": "sum", "field": "total"}]}, "统计单据总额")
        assert analysis["tools"][0]["result"]["totals"] == {"sum:total": 1000}
        save("versioned_edits_and_assistant", passed=True, analysis_thread=analysis["id"], expected_total=1000)

        cards = []
        for metric, field in [("count", None), ("sum", "total"), ("sum", "amount"), ("avg", "amount")]:
            cards.append(request("POST", "/stats/cards", json={"name": f"{metric}-{field or 'records'}",
                "table_id": table_id, "metric": metric, "metric_field": field}).json())
        assert client.post("/api/v1/stats/cards", json={"name": "第五张", "table_id": table_id}).status_code == 409
        values = {card["id"]: request("GET", f"/stats/cards/{card['id']}/result").json()["analysis"]["totals"] for card in cards}
        summary = request("GET", "/stats/summary?days=7").json()
        assert summary["current_tasks"] == 4 and summary["completed_count"] == 3
        assert summary["row_count"] == summary["new_rows"] == 6 and summary["review_pending"] == 0
        assert len(request("GET", table_url + "/export.json").json()) == 6
        assert len(list(csv.reader(StringIO(request("GET", table_url + "/export.csv").content.decode("utf-8-sig"))))) == 7
        exported = request("GET", table_url + "/export.xlsx").content
        book = load_workbook(BytesIO(exported))
        assert book.active.max_row == 7
        book.close()
        (root / "warehouse.xlsx").write_bytes(exported)
        before_table = request("GET", table_url).json()
        before_revisions = {row["id"]: request("GET", f"{table_url}/rows/{row['id']}/revisions").json() for row in before_table["rows"]}
        save("warehouse_dashboard_export", passed=True, row_count=6, card_count=4, card_totals=values, summary=summary)

    archive = create_business_backup(settings, root / "portable.dpbak")
    assert inspect_business_backup(archive)["credentials_included"] is False
    target = root / "restored"
    restored_settings = settings_for(target)
    restore_business_backup(restored_settings, archive, target)
    with TestClient(create_app(restored_settings)) as client:
        def get(path):
            response = client.get("/api/v1" + path)
            assert response.is_success, response.text
            return response
        assert get(table_url).json()["rows"] == before_table["rows"]
        assert [card["id"] for card in get("/stats/cards").json()["items"]] == [card["id"] for card in cards]
        for card in cards:
            assert get(f"/stats/cards/{card['id']}/result").json()["analysis"]["totals"] == values[card["id"]]
        for task_id, raw in originals.items():
            assert get(f"/tasks/{task_id}/file").content == raw
        for row_id, revisions in before_revisions.items():
            assert get(f"{table_url}/rows/{row_id}/revisions").json() == revisions
        assert get(f"/assistant/threads/{analysis['id']}").json()["tools"] == analysis["tools"]
        assert get(binding_url).json()["enabled"] is False
    with sqlite3.connect(target / "document-pipeline.db") as connection:
        assert connection.execute("PRAGMA integrity_check").fetchone() == ("ok",)
        revision = connection.execute("SELECT version_num FROM alembic_version").fetchone()[0]
    report.update(passed=True, offline_extraction_calls=model.calls, offline_chat_calls=OfflineConversation.calls)
    save("portable_restore", passed=True, schema=revision, original_count=len(originals), card_count=len(cards))
    print(json.dumps({"passed": True, "real_model_calls": 0, "root": str(root)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
