import argparse
from io import BytesIO
import json
import os
import signal
import subprocess
import sys
import tempfile
import time
from pathlib import Path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("file", type=Path)
    parser.add_argument("--expect-kind", choices=("invoice", "delivery"))
    parser.add_argument("--expect-number")
    parser.add_argument("--expect-rows", type=int)
    args = parser.parse_args()
    content_type = {
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".png": "image/png",
        ".pdf": "application/pdf",
    }[args.file.suffix.lower()]

    source_root = Path(__file__).resolve().parents[1] / "src"
    consumer_executable = Path(sys.executable).with_name("huey_consumer.exe")

    with tempfile.TemporaryDirectory() as temporary_directory:
        data_dir = Path(temporary_directory)
        os.environ["DOCUMENT_PIPELINE_DATA_DIR"] = str(data_dir)
        os.environ["DOCUMENT_PIPELINE_QUEUE_DB"] = str(data_dir / "queue.db")
        os.environ["PYTHONPATH"] = str(source_root)

        from fastapi.testclient import TestClient

        from document_pipeline_api.config import Settings
        from document_pipeline_api.main import create_app

        consumer = subprocess.Popen(
            [
                str(consumer_executable),
                "document_pipeline_api.worker.huey",
                "-w",
                "1",
                "-k",
                "thread",
                "-d",
                "0.01",
                "-m",
                "0.01",
                "-b",
                "1",
                "-q",
            ],
            cwd=Path.cwd(),
            env=os.environ.copy(),
            creationflags=subprocess.CREATE_NEW_PROCESS_GROUP,
        )
        try:
            with TestClient(create_app(Settings.local())) as client:
                with args.file.open("rb") as source:
                    response = client.post(
                        "/api/v1/tasks",
                        files={"file": (args.file.name, source, content_type)},
                        data={"template_mode": "smart"},
                    )
                response.raise_for_status()
                task_id = response.json()["id"]

                deadline = time.monotonic() + 90
                current = response.json()
                while time.monotonic() < deadline:
                    time.sleep(0.5)
                    current = client.get("/api/v1/tasks").json()[0]
                    if current["status"] in {"needs_review", "failed"}:
                        break

                result_response = client.get(f"/api/v1/tasks/{task_id}/result")
                review_response = None
                confirmation_response = None
                table_response = None
                export_response = None
                workbook_valid = False
                if result_response.status_code == 200:
                    extracted = result_response.json()
                    reviewed_result = extracted["result"]
                    original_seller = reviewed_result["seller_name"]
                    reviewed_result["seller_name"] = f"{original_seller or '未识别'}（人工复核）"
                    review_response = client.put(
                        f"/api/v1/tasks/{task_id}/review",
                        json={
                            "expected_version": extracted["review_version"],
                            "result": reviewed_result,
                        },
                    )
                    if review_response.status_code == 200:
                        confirmation_response = client.post(
                            f"/api/v1/tasks/{task_id}/confirm",
                            json={
                                "expected_review_version": review_response.json()[
                                    "review_version"
                                ],
                            },
                        )
                    if (
                        confirmation_response is not None
                        and confirmation_response.status_code == 200
                    ):
                        table_id = confirmation_response.json()["table_id"]
                        table_response = client.get(f"/api/v1/tables/{table_id}")
                        export_response = client.get(
                            f"/api/v1/tables/{table_id}/export.xlsx"
                        )
                        if export_response.status_code == 200:
                            from openpyxl import load_workbook

                            workbook = load_workbook(BytesIO(export_response.content))
                            workbook_valid = workbook.active.max_row >= 2
                final_task = client.get("/api/v1/tasks").json()[0]
                result_json = (
                    result_response.json()
                    if result_response.status_code == 200
                    else None
                )
                confirmation_json = (
                    confirmation_response.json()
                    if confirmation_response is not None
                    and confirmation_response.status_code == 200
                    else None
                )
                table_json = (
                    table_response.json()
                    if table_response is not None and table_response.status_code == 200
                    else None
                )
                print(
                    json.dumps(
                        {
                            "task_id": task_id,
                            "status": final_task["status"],
                            "consumer_exit_code": consumer.poll(),
                            "result_available": result_json is not None,
                            "document_kind": (
                                result_json["document_kind"] if result_json else None
                            ),
                            "document_number": (
                                result_json["result"]["document_number"]
                                if result_json
                                else None
                            ),
                            "review_version": (
                                review_response.json()["review_version"]
                                if review_response is not None
                                and review_response.status_code == 200
                                else None
                            ),
                            "model_result_preserved": (
                                review_response.json()["original_result"]["seller_name"]
                                == original_seller
                                if review_response is not None
                                and review_response.status_code == 200
                                else False
                            ),
                            "confirmed_table": (
                                confirmation_json["table_name"]
                                if confirmation_json
                                else None
                            ),
                            "table_rows": (
                                table_json["row_count"] if table_json else None
                            ),
                            "xlsx_valid": workbook_valid,
                        },
                        ensure_ascii=False,
                        indent=2,
                    )
                )
                if (
                    final_task["status"] != "completed"
                    or result_response.status_code != 200
                    or review_response is None
                    or review_response.status_code != 200
                    or confirmation_response is None
                    or confirmation_response.status_code != 200
                    or table_response is None
                    or table_response.status_code != 200
                    or not workbook_valid
                ):
                    raise SystemExit("真实队列闭环没有在期限内完成。")
                if (
                    args.expect_kind
                    and result_json["document_kind"] != args.expect_kind
                ):
                    raise SystemExit("文档分类与预期不一致。")
                if (
                    args.expect_number
                    and result_json["result"]["document_number"] != args.expect_number
                ):
                    raise SystemExit("单据号码与预期不一致。")
                if args.expect_rows and table_json["row_count"] != args.expect_rows:
                    raise SystemExit("展开后的事实行数与预期不一致。")
        finally:
            consumer.send_signal(signal.CTRL_BREAK_EVENT)
            consumer.wait(timeout=15)
            from document_pipeline_api.queue import huey

            huey.storage.close()
            time.sleep(0.5)


if __name__ == "__main__":
    main()
