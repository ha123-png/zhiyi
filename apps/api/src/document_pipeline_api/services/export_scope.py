"""Keep model input coverage visible without changing spreadsheet facts."""
import json
import re

from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font

from document_pipeline_api.schemas.input_scope import InputScope

SCOPE_KEY = "__zhiyi_input_scope"
SCOPE_LABEL = "知意输入范围（非业务字段）"
REVIEW_KEY = "__zhiyi_review_pending"
REVIEW_LABEL = "知意来源状态（非业务字段）"


def table_has_pending_reviews(session, table_id):
    from sqlalchemy import select
    from document_pipeline_api.models import DataRowRecord
    return session.scalar(select(DataRowRecord.id).where(DataRowRecord.table_id == table_id, DataRowRecord.review_pending.is_(True)).limit(1)) is not None


def imported_pending_review(value):
    if value in (None, "", False):
        return None
    if value is True or value in ("true", "True", "待核对"):
        return True
    raise ValueError("知意来源状态标记损坏，未导入。")


def table_has_input_scopes(session, table_id):
    from sqlalchemy import select
    from document_pipeline_api.models import DataRowRecord
    return session.scalar(select(DataRowRecord.id).where(DataRowRecord.table_id == table_id, DataRowRecord.input_scope_json.is_not(None)).limit(1)) is not None


def export_row_values(record):
    values = json.loads(record.row_json)
    if record.review_pending:
        if REVIEW_KEY in values:
            raise ValueError("业务字段与知意来源状态保留字段冲突，未导出。")
        values[REVIEW_KEY] = True
    if record.input_scope_json:
        if SCOPE_KEY in values:
            raise ValueError("业务字段与知意输入范围保留字段冲突，未导出。")
        values[SCOPE_KEY] = InputScope.model_validate_json(record.input_scope_json).model_dump(mode="json")
    return values


def csv_scope_value(values):
    scope = values.get(SCOPE_KEY)
    return json.dumps(scope, ensure_ascii=False) if scope else ""


class WorkbookInputScopes:
    def __init__(self):
        self.sources = {}
        self.cells = []

    def record(self, sheet, row_number, record, keys, header_keys):
        if record.review_pending:
            if REVIEW_LABEL in [cell.value for cell in sheet[1][:len(keys)]]:
                raise ValueError("业务字段与知意来源状态说明列冲突，未导出。")
            column = len(keys) + 1
            sheet.cell(1, column, REVIEW_LABEL)
            sheet.column_dimensions[sheet.cell(1, column).column_letter].width = 28
            cell = sheet.cell(row_number, column, "待核对")
            cell.comment = Comment("来源任务仍有待处理事项。编辑数据不等于重新执行校验；来源已移除或本表由导入/合并生成时，这是保留的来源状态。其他行没有此标记也不代表经过人工核验。", "知意")
        if not record.input_scope_json:
            return
        scope = InputScope.model_validate_json(record.input_scope_json)
        values = json.loads(record.row_json)
        source_key = (record.task_id or values.get("__row_group") or record.id, record.input_scope_json)
        if source_key not in self.sources:
            self.sources[source_key] = (len(self.sources) + 1, str(values.get("source_filename") or "未记录原文件名"), scope)
        source_id, _, _ = self.sources[source_key]
        # Header fields may be merged vertically later; place the reference in
        # an item cell whenever one exists so every content row keeps its link.
        column = next((index for index, key in enumerate(keys, 1) if key not in header_keys), 1)
        self.cells.append((sheet.cell(row_number, column), source_id, scope.coverage))

    def finish(self, workbook):
        if not self.sources:
            return
        title = "知意输入范围"
        index = 2
        while title.casefold() in {name.casefold() for name in workbook.sheetnames}:
            title = f"知意输入范围-{index}"
            index += 1
        sheet = workbook.create_sheet(title)
        sheet.append(["来源编号", "原上传文件名", "输入覆盖", "类别", "原始范围或说明", "原因"])
        sheet.append([None, None, None, "使用说明", "业务表单元格批注中的来源编号对应本页。范围记录模型当时收到什么，不代表提取正确率；局部读取不能代表全文或全部记录。没有批注的旧行未记录范围，不代表完整读取。", None])
        reasons = {"input_budget": "达到当时设置的读取上限", "images_disabled": "未开启图片读取", "unsupported_image": "暂不支持的图片", "unsupported_structure": "暂不支持的文档部分"}
        for source_id, filename, scope in self.sources.values():
            coverage = "局部读取" if scope.coverage == "partial" else "完整提供（不保证提取正确）"
            prefix = [source_id, _safe_text(filename), coverage]
            sheet.append([*prefix, "规则", scope.rule, None])
            for location in scope.selected:
                sheet.append([*prefix, "已提供", _safe_text(location.label()), None])
            for omitted in scope.omitted:
                sheet.append([*prefix, "未提供", _safe_text(omitted.location.label()), reasons[omitted.reason]])
            for note in scope.notes:
                # Excel cells have a hard text length limit. Split notes into
                # explicit continuation rows, never let openpyxl truncate them.
                for offset in range(0, len(note), 30000):
                    sheet.append([*prefix, "说明" if offset == 0 else "说明（续）", _safe_text(note[offset:offset + 30000]), None])
            encoded = scope.model_dump_json()
            for index, offset in enumerate(range(0, len(encoded), 30000), 1):
                sheet.append([source_id, None, None, "内部范围记录 v1", "json:" + encoded[offset:offset + 30000], index])
                sheet.row_dimensions[sheet.max_row].hidden = True
        sheet.freeze_panes = "A3"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="065F46")
        for column, width in {"A": 12, "B": 32, "C": 28, "D": 16, "E": 72, "F": 24}.items():
            sheet.column_dimensions[column].width = width
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell, source_id, coverage in self.cells:
            label = "局部读取，不能代表全文或全部记录" if coverage == "partial" else "完整提供，不保证提取正确"
            cell.comment = Comment(f"{label}。来源编号 {source_id}，实际提供和遗漏的范围见「{title}」工作表。", "知意")


def _safe_text(value):
    return "'" + value if value.startswith(("=", "+", "-", "@")) else value


def read_workbook_scopes(workbook, business_sheet):
    """Restore only explicitly linked metadata, never infer complete coverage."""
    result = {}
    cache = {}
    for row in business_sheet.iter_rows(min_row=2):
        for cell in row:
            comment = cell.comment
            if not comment or comment.author != "知意":
                continue
            match = re.search(r"来源编号 (\d+)，实际提供和遗漏的范围见「(.+)」工作表。", comment.text)
            if not match:
                continue
            source_id, title = int(match[1]), match[2]
            if title not in cache:
                if title not in workbook.sheetnames:
                    raise ValueError("Excel 的知意输入范围说明页已丢失，无法保留来源范围。")
                sources = {}
                for record in workbook[title].iter_rows(min_row=3, values_only=True):
                    if len(record) >= 6 and record[3] == "内部范围记录 v1":
                        identifier, text, part = record[0], record[4], record[5]
                        if not isinstance(identifier, int) or not isinstance(part, int) or not isinstance(text, str) or not text.startswith("json:"):
                            raise ValueError("Excel 输入范围记录损坏。")
                        parts = sources.setdefault(identifier, {})
                        if part in parts:
                            raise ValueError("Excel 输入范围记录包含重复片段。")
                        parts[part] = text[5:]
                decoded = {}
                for identifier, parts in sources.items():
                    if sorted(parts) != list(range(1, len(parts) + 1)):
                        raise ValueError("Excel 输入范围记录缺少片段。")
                    decoded[identifier] = InputScope.model_validate_json("".join(parts[index] for index in sorted(parts))).model_dump_json()
                cache[title] = decoded
            if source_id not in cache[title]:
                raise ValueError("Excel 的来源批注找不到对应范围记录。")
            value = cache[title][source_id]
            if cell.row in result and result[cell.row] != value:
                raise ValueError("Excel 同一数据行关联了冲突的输入范围。")
            result[cell.row] = value
    return result
