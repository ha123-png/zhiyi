"""文件格式归一化：图片转换与 Office/文本提取。

首版边界：
- 图片：webp/bmp/tiff/gif 等 Pillow 可读格式转 PNG（保留原文件作证据，转换产物进 rendered/）。
- Office/文本：txt/md 直接读取；docx 解析 word/document.xml；xlsx 用 openpyxl 读单元格。
  提取出的文本按行分页并渲染成白底黑字 PNG 页，交给现有图像识别链路。
  这与"Word 转 PDF 后再识别"相比是可靠的纯文字模式：不依赖 Office/LibreOffice。
"""

from __future__ import annotations

import os
from io import BytesIO
from pathlib import Path
import zipfile
import warnings
from xml.etree import ElementTree

from PIL import Image, ImageDraw, ImageFont

# 需要归一化的图片类型（jpeg/png 是模型原生可读，无需转换）
CONVERTIBLE_IMAGE_TYPES = {
    "image/webp": ".webp",
    "image/bmp": ".bmp",
    "image/tiff": ".tiff",
    "image/gif": ".gif",
}
RASTER_IMAGE_TYPES = {
    "image/jpeg": ".jpg",
    "image/png": ".png",
    **CONVERTIBLE_IMAGE_TYPES,
}
_PIL_FORMATS_BY_CONTENT_TYPE = {
    "image/jpeg": {"JPEG"},
    "image/png": {"PNG"},
    "image/webp": {"WEBP"},
    "image/bmp": {"BMP"},
    "image/tiff": {"TIFF"},
    "image/gif": {"GIF"},
}

# 文本/Office 类型（上传后提取文字并渲染为图片页）
TEXT_CONTENT_TYPES = {
    "text/plain": ".txt",
    "text/markdown": ".md",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
}

MAX_LINES_PER_PAGE = 45
_PAGE_WIDTH = 1240
_LINE_HEIGHT = 34
_FONT_SIZE = 26
_MARGIN = 48


class UnsupportedTextFileError(ValueError):
    pass


class UnsupportedImageError(ValueError):
    pass


def verify_image(path: Path) -> None:
    """验证 Pillow 能否读取该图片（上传准入用，不落盘）。"""
    inspect_image_frame_count(path)


def inspect_image_frame_count(
    path: Path,
    *,
    expected_content_type: str | None = None,
    max_frames: int | None = None,
    max_total_pixels: int | None = None,
) -> int:
    """验证真实格式并逐帧解码；返回可安全处理的总帧数。"""
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(path) as source:
                actual_format = (source.format or "").upper()
                expected_formats = _PIL_FORMATS_BY_CONTENT_TYPE.get(
                    expected_content_type or ""
                )
                if expected_formats is not None and actual_format not in expected_formats:
                    raise UnsupportedImageError(
                        "图片实际格式与上传时声明的类型不一致，请勿修改扩展名后重试。"
                    )
                frame_count = int(getattr(source, "n_frames", 1))
                if frame_count < 1:
                    raise UnsupportedImageError("图片没有可读取的页面。")
                if max_frames is not None and frame_count > max_frames:
                    raise UnsupportedImageError(
                        f"图片共 {frame_count} 页，超过当前 {max_frames} 页处理上限。"
                    )
                total_pixels = 0
                # verify() 只保证当前帧；逐帧 load 才能发现后续帧截断/损坏。
                for index in range(frame_count):
                    source.seek(index)
                    width, height = source.size
                    if width < 1 or height < 1:
                        raise UnsupportedImageError("图片包含尺寸无效的页面。")
                    total_pixels += width * height
                    if max_total_pixels is not None and total_pixels > max_total_pixels:
                        raise UnsupportedImageError(
                            f"图片总像素超过当前 {max_total_pixels} 像素处理上限。"
                        )
                    source.load()
                return frame_count
    except UnsupportedImageError:
        raise
    except (Image.DecompressionBombError, Image.DecompressionBombWarning) as error:
        raise UnsupportedImageError("图片像素尺寸过大，无法安全处理。") from error
    except Exception as error:
        raise UnsupportedImageError("图片文件损坏或无法读取，请更换后重试。") from error


def _find_cjk_font() -> Path:
    windows_dir = Path(os.environ.get("WINDIR", "C:/Windows"))
    candidates = [
        windows_dir / "Fonts" / "msyh.ttc",   # 微软雅黑
        windows_dir / "Fonts" / "simsun.ttc",  # 宋体
        windows_dir / "Fonts" / "simhei.ttf",  # 黑体
        windows_dir / "Fonts" / "msyhbd.ttc",
    ]
    for candidate in candidates:
        if candidate.is_file():
            return candidate
    raise UnsupportedTextFileError(
        "系统缺少可用于渲染的中文字体（微软雅黑/宋体），无法转换文本或 Office 文档。"
    )


def normalize_image(src: Path, dst: Path) -> None:
    """把单张图片的当前首帧转成白底 RGB PNG（兼容旧调用）。"""
    dst.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(src) as source:
        _image_frame_to_rgb(source).save(dst, "PNG")


def render_image_frames(
    src: Path,
    out_dir: Path,
    task_id: str,
    *,
    max_frames: int,
    expected_content_type: str | None = None,
    max_total_pixels: int | None = None,
) -> list[Path]:
    """把图片的全部帧按原顺序渲染为独立 PNG 页。"""
    frame_count = inspect_image_frame_count(
        src,
        expected_content_type=expected_content_type,
        max_frames=max_frames,
        max_total_pixels=max_total_pixels,
    )
    out_dir.mkdir(parents=True, exist_ok=True)
    rendered: list[Path] = []
    try:
        with Image.open(src) as source:
            for index in range(frame_count):
                source.seek(index)
                target = out_dir / f"{task_id}-image-{index + 1}.png"
                _image_frame_to_rgb(source).save(target, "PNG")
                rendered.append(target)
    except Exception as error:
        for target in rendered:
            target.unlink(missing_ok=True)
        raise UnsupportedImageError("图片页面转换失败，请更换文件后重试。") from error
    return rendered


def _image_frame_to_rgb(source: Image.Image) -> Image.Image:
    if source.mode in ("RGBA", "LA", "P"):
        rgba = source.convert("RGBA")
        background = Image.new("RGB", rgba.size, (255, 255, 255))
        background.paste(rgba, mask=rgba.getchannel("A"))
        return background
    return source.convert("RGB")


def _validate_zip_budget(path: Path, *, max_uncompressed_bytes: int | None) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            total = 0
            for member in archive.infolist():
                if member.flag_bits & 0x1:
                    raise UnsupportedTextFileError("Office 文档包含加密内容，当前无法处理。")
                total += member.file_size
                if max_uncompressed_bytes is not None and total > max_uncompressed_bytes:
                    raise UnsupportedTextFileError(
                        f"Office 文档解压后超过当前 {max_uncompressed_bytes} 字节处理上限。"
                    )
    except UnsupportedTextFileError:
        raise
    except (zipfile.BadZipFile, OSError) as error:
        raise UnsupportedTextFileError("Office 文档损坏或不是有效的压缩文档。") from error


def _extract_docx_text(path: Path, *, max_uncompressed_bytes: int | None) -> str:
    _validate_zip_budget(path, max_uncompressed_bytes=max_uncompressed_bytes)
    try:
        with zipfile.ZipFile(path) as archive:
            xml_bytes = archive.read("word/document.xml")
    except (zipfile.BadZipFile, KeyError, OSError) as error:
        raise UnsupportedTextFileError("Word 文档损坏或不是有效的 .docx 文件。") from error
    root = ElementTree.fromstring(xml_bytes)
    paragraphs: list[str] = []
    for element in root.iter():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = "".join(
                node.text or ""
                for node in element.iter()
                if node.tag.rsplit("}", 1)[-1] == "t"
            )
            stripped = text.strip()
            if stripped:
                paragraphs.append(stripped)
    return "\n".join(paragraphs)


# docx 结构化预览命名空间与块类型
_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def extract_docx_blocks(
    path: Path,
    *,
    max_uncompressed_bytes: int | None,
    max_images: int | None = None,
    max_rows: int | None = None,
    max_columns: int | None = None,
) -> list[dict[str, object]]:
    """把 docx 按文档结构读成有序块（标题/段落/列表/表格/按位图片），供只读预览。

    保留标题层级、表格网格和内嵌图片在原文档中的相对位置，接近原貌但不依赖
    Office/LibreOffice 转 PDF。图片以 ``image_index`` 引用 media 列表（1 基），
    由图片接口按同一顺序服务。
    """
    _validate_zip_budget(path, max_uncompressed_bytes=max_uncompressed_bytes)
    try:
        with zipfile.ZipFile(path) as archive:
            document_bytes = archive.read("word/document.xml")
            rels_bytes = archive.read("word/_rels/document.xml.rels")
            media_names = sorted(
                info.filename
                for info in archive.infolist()
                if info.filename.startswith("word/media/") and not info.is_dir()
            )
    except (zipfile.BadZipFile, KeyError, OSError) as error:
        raise UnsupportedTextFileError("Word 文档损坏或不是有效的 .docx 文件。") from error
    if max_images is not None:
        media_names = media_names[:max_images]
    media_index = {name.split("/")[-1]: index for index, name in enumerate(media_names)}
    rels = _docx_rels_map(rels_bytes)
    blocks: list[dict[str, object]] = []
    try:
        root = ElementTree.fromstring(document_bytes)
    except ElementTree.ParseError as error:
        raise UnsupportedTextFileError("Word 文档结构损坏，无法读取正文。") from error
    body = root.find(f"{{{_W_NS}}}body")
    if body is not None:
        for child in body:
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                _docx_paragraph_blocks(child, rels, media_index, blocks)
            elif tag == "tbl":
                _docx_table_block(child, blocks, max_rows=max_rows, max_columns=max_columns)
    return blocks


def _docx_rels_map(rels_bytes: bytes) -> dict[str, str]:
    """解析 word/_rels/document.xml.rels：rId -> 媒体目标文件名。

    ``Id``/``Target`` 是无命名空间属性，直接按字面读取。
    """
    try:
        root = ElementTree.fromstring(rels_bytes)
    except ElementTree.ParseError:
        return {}
    mapping: dict[str, str] = {}
    for rel in root:
        rid = rel.get("Id")
        target = rel.get("Target") or ""
        if rid:
            mapping[rid] = target
    return mapping


def _docx_heading_level(p_pr) -> int:
    style = p_pr.find(f"{{{_W_NS}}}pStyle")
    if style is None:
        return 0
    val = (style.get(f"{{{_W_NS}}}val") or "").strip()
    lower = val.lower()
    digits = "".join(ch for ch in lower if ch.isdigit())
    if ("heading" in lower or "标题" in val) and digits:
        return int(digits)
    return 0


def _docx_paragraph_blocks(
    p,
    rels: dict[str, str],
    media_index: dict[str, int],
    blocks: list[dict[str, object]],
) -> None:
    p_pr = p.find(f"{{{_W_NS}}}pPr")
    heading_level = _docx_heading_level(p_pr) if p_pr is not None else 0
    # 段落内按出现顺序收集图片引用（rId -> media）
    image_refs: list[int] = []
    for drawing in p.iter(f"{{{_W_NS}}}drawing"):
        for blip in drawing.iter(f"{{{_A_NS}}}blip"):
            embed = blip.get(f"{{{_R_NS}}}embed")
            if not embed:
                continue
            target = (rels.get(embed) or "").split("/")[-1]
            if target in media_index:
                image_refs.append(media_index[target] + 1)
    text = "".join(node.text or "" for node in p.iter(f"{{{_W_NS}}}t")).strip()
    if heading_level:
        if text:
            blocks.append({"type": "heading", "level": heading_level, "text": text})
        return
    # 纯文本段：区分列表项与普通段落
    is_list = p_pr is not None and p_pr.find(f"{{{_W_NS}}}numPr") is not None
    if text and not image_refs:
        if is_list:
            blocks.append({"type": "list", "text": text})
        else:
            blocks.append({"type": "paragraph", "text": text})
    else:
        if text:
            blocks.append({"type": "paragraph", "text": text})
        for image_index in image_refs:
            blocks.append(
                {
                    "type": "image",
                    "image_index": image_index,
                    "caption": f"图片 {image_index}",
                }
            )


def _docx_table_block(
    tbl,
    blocks: list[dict[str, object]],
    *,
    max_rows: int | None,
    max_columns: int | None,
) -> None:
    rows: list[list[str]] = []
    for tr in tbl.iter(f"{{{_W_NS}}}tr"):
        row: list[str] = []
        for tc in tr.iter(f"{{{_W_NS}}}tc"):
            cell = " ".join(
                "".join(node.text or "" for node in p.iter(f"{{{_W_NS}}}t"))
                for p in tc.iter(f"{{{_W_NS}}}p")
            ).strip()
            row.append(cell)
        if max_rows is not None and len(rows) > max_rows:
            raise UnsupportedTextFileError(f"Word 表格超过当前 {max_rows} 行处理上限。")
        if max_columns is not None and len(row) > max_columns:
            raise UnsupportedTextFileError(f"Word 表格超过当前 {max_columns} 列处理上限。")
        rows.append(row)
    blocks.append({"type": "table", "rows": rows})


def extract_docx_images(
    path: Path,
    out_dir: Path,
    task_id: str,
    *,
    max_uncompressed_bytes: int | None,
    max_images: int | None = None,
    max_total_pixels: int | None = None,
) -> list[Path]:
    """把 docx 内嵌位图按原顺序归一化为白底 RGB PNG 页，供模型识别。

    非位图媒体（emf/wmf 等）与损坏图片跳过，不影响正文识别。
    """
    _validate_zip_budget(path, max_uncompressed_bytes=max_uncompressed_bytes)
    try:
        with zipfile.ZipFile(path) as archive:
            media_names = sorted(
                info.filename
                for info in archive.infolist()
                if info.filename.startswith("word/media/") and not info.is_dir()
            )
            if max_images is not None:
                media_names = media_names[:max_images]
            payloads = [(name, archive.read(name)) for name in media_names]
    except (zipfile.BadZipFile, OSError) as error:
        raise UnsupportedTextFileError("Word 文档损坏或不是有效的 .docx 文件。") from error
    out_dir.mkdir(parents=True, exist_ok=True)
    rendered: list[Path] = []
    for index, (name, data) in enumerate(payloads):
        if name.lower().rsplit(".", 1)[-1] in {"emf", "wmf"}:
            continue
        try:
            with Image.open(BytesIO(data)) as source:
                frame = _image_frame_to_rgb(source)
                if max_total_pixels is not None:
                    frame = _fit_total_pixels(frame, max_total_pixels)
                target = out_dir / f"{task_id}-docx-image-{index + 1}.png"
                frame.save(target, "PNG")
                rendered.append(target)
        except Exception:
            continue
    return rendered


def _fit_total_pixels(image: Image.Image, limit: int) -> Image.Image:
    """总像素超限时等比缩小，避免超大内嵌图撑爆模型输入。"""
    width, height = image.size
    if width * height <= limit:
        return image
    ratio = (limit / (width * height)) ** 0.5
    return image.resize((max(1, int(width * ratio)), max(1, int(height * ratio))), Image.LANCZOS)


def _extract_xlsx_text(
    path: Path,
    *,
    max_uncompressed_bytes: int | None,
    max_rows: int | None,
    max_columns: int | None,
) -> str:
    _validate_zip_budget(path, max_uncompressed_bytes=max_uncompressed_bytes)
    try:
        from openpyxl import load_workbook
    except ImportError as error:  # pragma: no cover
        raise UnsupportedTextFileError("当前环境缺少 Excel 读取组件。") from error
    source = None
    try:
        # Upload validation runs before the staging file is renamed to ``.xlsx``.
        # Passing that ``.<uuid>.uploading`` path directly makes openpyxl reject a
        # valid workbook solely because of its temporary suffix.  A binary stream
        # keeps the same parser and ZIP checks without coupling validity to a name.
        source = path.open("rb")
        workbook = load_workbook(source, read_only=True, data_only=True)
    except Exception as error:
        if source is not None:
            source.close()
        raise UnsupportedTextFileError("Excel 文件损坏或不是有效的 .xlsx 文件。") from error
    lines: list[str] = []
    row_count = 0
    try:
        for sheet in workbook.worksheets:
            if workbook.worksheets.index(sheet) > 0:
                lines.append("")
            lines.append(f"【工作表】{sheet.title}")
            for row in sheet.iter_rows():
                row_count += 1
                if max_rows is not None and row_count > max_rows:
                    raise UnsupportedTextFileError(
                        f"Excel 文档超过当前 {max_rows} 行处理上限。"
                    )
                if max_columns is not None and len(row) > max_columns:
                    raise UnsupportedTextFileError(
                        f"Excel 文档超过当前 {max_columns} 列处理上限。"
                    )
                cells = [
                    str(cell.value) if cell.value is not None else ""
                    for cell in row
                ]
                if any(cells):
                    lines.append("\t".join(cells))
    finally:
        workbook.close()
        source.close()
    return "\n".join(lines)


def _excel_column_name(index: int) -> str:
    """0 基索引转 Excel 列名（0->A, 25->Z, 26->AA）。"""
    name = ""
    index += 1
    while index:
        index, remainder = divmod(index - 1, 26)
        name = chr(65 + remainder) + name
    return name


def extract_xlsx_sheets(
    path: Path,
    *,
    max_uncompressed_bytes: int | None = None,
    max_rows: int | None = None,
    max_columns: int | None = None,
) -> list[dict[str, object]]:
    """把 xlsx 按工作表读成结构化表格（列名 + 单元格行），供只读表格预览。"""
    _validate_zip_budget(path, max_uncompressed_bytes=max_uncompressed_bytes)
    try:
        from openpyxl import load_workbook
    except ImportError as error:  # pragma: no cover
        raise UnsupportedTextFileError("当前环境缺少 Excel 读取组件。") from error
    source = None
    try:
        source = path.open("rb")
        workbook = load_workbook(source, read_only=True, data_only=True)
    except Exception as error:
        if source is not None:
            source.close()
        raise UnsupportedTextFileError("Excel 文件损坏或不是有效的 .xlsx 文件。") from error
    sheets: list[dict[str, object]] = []
    try:
        for sheet in workbook.worksheets:
            rows: list[list[str]] = []
            for row_index, row in enumerate(sheet.iter_rows(), start=1):
                if max_rows is not None and row_index > max_rows:
                    raise UnsupportedTextFileError(
                        f"Excel 文档超过当前 {max_rows} 行处理上限。"
                    )
                if max_columns is not None and len(row) > max_columns:
                    raise UnsupportedTextFileError(
                        f"Excel 文档超过当前 {max_columns} 列处理上限。"
                    )
                rows.append(
                    [
                        str(cell.value) if cell.value is not None else ""
                        for cell in row
                    ]
                )
            column_count = max((len(row) for row in rows), default=0)
            columns = [_excel_column_name(i) for i in range(column_count)]
            sheets.append({"name": sheet.title, "columns": columns, "rows": rows})
    finally:
        workbook.close()
        source.close()
    return sheets


def extract_text(
    path: Path,
    content_type: str,
    *,
    max_uncompressed_bytes: int | None = None,
    max_rows: int | None = None,
    max_columns: int | None = None,
) -> str:
    """按类型提取纯文本；文本将按行分页渲染为图片。"""
    if content_type in {"text/plain", "text/markdown"}:
        raw = path.read_bytes()
        for encoding in ("utf-8", "gbk"):
            try:
                return raw.decode(encoding)
            except UnicodeDecodeError:
                continue
        raise UnsupportedTextFileError("文本文件编码无法识别，仅支持 UTF-8 或 GBK。")
    if content_type.endswith(".wordprocessingml.document"):
        return _extract_docx_text(
            path, max_uncompressed_bytes=max_uncompressed_bytes
        )
    if content_type.endswith(".spreadsheetml.sheet"):
        return _extract_xlsx_text(
            path,
            max_uncompressed_bytes=max_uncompressed_bytes,
            max_rows=max_rows,
            max_columns=max_columns,
        )
    raise UnsupportedTextFileError("不支持的文本文件类型。")


def split_text_pages(text: str, *, max_pages: int) -> list[str]:
    """按行分页；超过 max_pages 页报错（与 PDF 页数上限语义一致）。"""
    lines = [line for line in text.splitlines() if line.strip()]
    if not lines:
        raise UnsupportedTextFileError("文档没有可提取的正文内容。")
    page_count = max(1, (len(lines) + MAX_LINES_PER_PAGE - 1) // MAX_LINES_PER_PAGE)
    if page_count > max_pages:
        raise UnsupportedTextFileError(
            f"文档共约 {page_count} 页文字，超过当前 {max_pages} 页处理上限。"
        )
    pages: list[str] = []
    for index in range(0, len(lines), MAX_LINES_PER_PAGE):
        pages.append("\n".join(lines[index : index + MAX_LINES_PER_PAGE]))
    return pages


def render_text_pages(text_pages: list[str], out_dir: Path, task_id: str) -> list[Path]:
    """把文本页渲染成白底黑字 PNG 页，返回按页序排列的图片路径。"""
    font_path = _find_cjk_font()
    font = ImageFont.truetype(str(font_path), _FONT_SIZE)
    out_dir.mkdir(parents=True, exist_ok=True)
    rendered: list[Path] = []
    for index, page_text in enumerate(text_pages, start=1):
        target = out_dir / f"{task_id}-text-{index}.png"
        line_count = len(page_text.splitlines())
        height = max(400, _MARGIN * 2 + _LINE_HEIGHT * line_count)
        image = Image.new("RGB", (_PAGE_WIDTH, height), (255, 255, 255))
        draw = ImageDraw.Draw(image)
        y = _MARGIN
        for line in page_text.splitlines():
            draw.text((_MARGIN, y), line, fill=(0, 0, 0), font=font)
            y += _LINE_HEIGHT
        image.save(target, "PNG")
        rendered.append(target)
    return rendered
