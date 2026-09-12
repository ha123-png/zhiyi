from pathlib import Path
from io import BytesIO

from fastapi.testclient import TestClient
import pypdfium2 as pdfium
import pytest
from openpyxl import load_workbook

from document_pipeline_api.config import Settings
from document_pipeline_api.main import create_app
from document_pipeline_api.models import TaskRecord
from document_pipeline_api.schemas.templates import TemplateMatchDecision
from document_pipeline_api.services.extraction import process_task


@pytest.fixture
def client(tmp_path):
    settings = Settings(database_url=f"sqlite:///{tmp_path / 'scoped.db'}", storage_dir=tmp_path / "uploads", max_pdf_pages=4, max_active_pages=5)
    with TestClient(create_app(settings)) as client:
        yield client


def template(client, *, complete=True, name="局部概览"):
    response = client.post("/api/v1/templates", json={"name": name,
        "fields": [{"key": "title", "label": "标题", "section": "header"}],
        "behavior": {"requires_complete_input": complete}})
    assert response.status_code == 201
    return response.json()


def limited(client, enabled=True):
    response = client.put("/api/v1/system/settings", json={"office_convert": True, "image_convert": True, "allow_limited_input": enabled, "input_text_limit": 2048, "input_page_limit": 4})
    assert response.status_code == 200


LONG_TEXT = "开头文件信息\n" + "\n".join(f"第{i}行：这是一段测试正文，不是完整摘要。" for i in range(1000)) + "\n末尾凭证编号 FINAL-8765"


class TextModel:
    model_name = "scope-test"

    def __init__(self, matches=None):
        self.calls = []
        self.matches = matches

    def complete_text(self, prompt, result_type):
        self.calls.append(prompt)
        assert "开头文件信息" in prompt
        assert "末尾凭证编号 FINAL-8765" not in prompt
        assert "不能声称全文摘要或全部明细" in prompt
        if result_type is TemplateMatchDecision:
            return result_type(outcome="ambiguous", template_ids=self.matches)
        return result_type.model_validate({"header": {"title": "FINAL-8765"}, "items": []})


def run(client, task_id, model):
    with client.app.state.session_factory() as session:
        return process_task(session, client.app.state.settings, task_id, client=model)


def test_default_provides_complete_input_without_context_guess(client):
    response = client.post("/api/v1/tasks", files={"file": ("large.txt", LONG_TEXT, "text/plain")})
    assert response.status_code == 201
    assert response.json()["planned_scope"]["coverage"] == "complete"
    assert client.get(f"/api/v1/tasks/{response.json()['id']}/file").content.decode() == LONG_TEXT


def test_global_scope_policy_snapshot_review_and_original_survive(client):
    selected = template(client)
    limited(client)
    uploaded = client.post("/api/v1/tasks", files={"file": ("large.md", LONG_TEXT, "text/markdown")}, data={"template_id": selected["id"]})
    assert uploaded.status_code == 201, uploaded.text
    task = uploaded.json()
    assert task["planned_scope"]["coverage"] == "partial"
    # Changes after upload must not silently change an old task's input policy or template.
    limited(client, False)
    changed = client.put(f"/api/v1/templates/{selected['id']}", json={"expected_version": 1, "name": selected["name"], "fields": selected["fields"], "behavior": {"requires_complete_input": False}})
    assert changed.status_code == 200
    model = TextModel()
    # The uploaded global policy is authoritative; a legacy template flag
    # must not add another scope-consent gate.
    result = run(client, task["id"], model)
    assert result is not None and result.input_scope.coverage == "partial"
    assert len(model.calls) == 1
    assert all(item.page_number is None for item in result.evidence)
    assert any(issue.code == "partial_input" for issue in result.validation_issues)
    review = client.put(f"/api/v1/tasks/{task['id']}/review", json={"expected_version": 0, "result": result.result.model_dump()})
    assert review.status_code == 200, review.text
    assert review.json()["input_scope"] == result.input_scope.model_dump()
    assert any(issue["code"] == "partial_input" for issue in review.json()["validation_issues"])
    with client.app.state.session_factory() as session:
        stored = session.get(TaskRecord, task["id"])
        assert stored.template_version == 1
        assert stored.status == "needs_review"
        assert (client.app.state.settings.storage_dir / stored.storage_path).read_text(encoding="utf-8") == LONG_TEXT
    confirmed = client.post(f"/api/v1/tasks/{task['id']}/confirm", json={"expected_review_version": 1})
    assert confirmed.status_code == 200, confirmed.text
    table_id = confirmed.json()["table_id"]
    expected_scope = result.input_scope.model_dump()
    row = client.get(f"/api/v1/tables/{table_id}").json()["rows"][0]
    assert row["input_scope"] == expected_scope
    second = client.post("/api/v1/tables", json={"name": "空表", "template_key": selected["id"]})
    assert second.status_code == 201, second.text
    merged = client.post("/api/v1/tables/merge", json={"name": "合并保留范围", "table_ids": [table_id, second.json()["id"]]})
    assert merged.status_code == 201, merged.text
    merged_row = client.get(f"/api/v1/tables/{merged.json()['id']}").json()["rows"][0]
    assert merged_row["task_id"] is None
    assert merged_row["input_scope"] == expected_scope
    assert client.post("/api/v1/system/admin/clear-history", json={"confirm_text": "清除历史"}).status_code == 200
    remaining = client.get(f"/api/v1/tables/{table_id}").json()["rows"][0]
    assert remaining["task_id"] is None
    assert remaining["input_scope"] == expected_scope
    exported = client.get(f"/api/v1/tables/{table_id}/export.xlsx")
    assert exported.status_code == 200, exported.text
    workbook = load_workbook(BytesIO(exported.content))
    assert "知意输入范围" in workbook.sheetnames
    assert "局部读取" in workbook.worksheets[0].cell(2, 1).comment.text
    recorded = list(workbook["知意输入范围"].values)
    assert any(row[3] == "未提供" for row in recorded)
    for source_range in result.input_scope.selected:
        assert any(row[4] == source_range.label() for row in recorded)
    xlsx_imported = client.post("/api/v1/tables/import-new", files={"file": ("范围随行.xlsx", exported.content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert xlsx_imported.status_code == 201, xlsx_imported.text
    xlsx_row = client.get(f"/api/v1/tables/{xlsx_imported.json()['id']}").json()["rows"][0]
    assert xlsx_row["input_scope"]["selected"] == expected_scope["selected"]
    assert xlsx_row["input_scope"]["coverage"] == "partial"
    assert xlsx_row["task_id"] is None
    json_export = client.get(f"/api/v1/tables/{table_id}/export.json")
    assert json_export.status_code == 200
    assert json_export.json()[0]["__zhiyi_input_scope"] == expected_scope
    csv_export = client.get(f"/api/v1/tables/{table_id}/export.csv")
    assert csv_export.status_code == 200
    assert "知意输入范围（非业务字段）" in csv_export.text
    imported = client.post("/api/v1/tables/import-new", files={"file": ("范围随行.csv", csv_export.content, "text/csv")})
    assert imported.status_code == 201, imported.text
    imported_table = client.get(f"/api/v1/tables/{imported.json()['id']}").json()
    imported_row = imported_table["rows"][0]
    assert imported_row["task_id"] is None
    assert imported_row["input_scope"]["selected"] == expected_scope["selected"]
    assert imported_row["input_scope"]["coverage"] == "partial"
    assert any("未重新读取或核验原件" in note for note in imported_row["input_scope"]["notes"])
    assert all(column["label"] != "知意输入范围（非业务字段）" for column in imported_table["columns"])
    from document_pipeline_api.services.integration_queries import export_data_table_file
    import json
    import csv
    for extension in ("json", "csv"):
        path = client.app.state.settings.storage_dir.parent / f"mcp-range.{extension}"
        with client.app.state.session_factory() as session:
            assert export_data_table_file(session, table_id, path, extension) == 1
        if extension == "json":
            assert json.loads(path.read_text(encoding="utf-8"))[0]["__zhiyi_input_scope"] == expected_scope
        else:
            with path.open(encoding="utf-8-sig", newline="") as stream:
                exported_rows = list(csv.DictReader(stream))
            assert json.loads(exported_rows[0]["知意输入范围（非业务字段）"]) == expected_scope


def test_smart_ambiguity_records_match_range_without_formal_extraction(client):
    first = template(client, complete=False, name="报告概览")
    second = template(client, complete=False, name="笔记概览")
    limited(client)
    task = client.post("/api/v1/tasks", files={"file": ("smart.txt", LONG_TEXT, "text/plain")}).json()
    model = TextModel([first["id"], second["id"]])
    assert run(client, task["id"], model) is None
    waiting = client.get("/api/v1/tasks").json()[0]
    assert waiting["pending_reason"] == "template"
    assert waiting["match_scope"]["coverage"] == "partial"
    assert client.get(f"/api/v1/tasks/{task['id']}/result").status_code == 404
    assert client.post(f"/api/v1/tasks/{task['id']}/input-scope").status_code == 409
    assert client.post(f"/api/v1/tasks/{task['id']}/template", json={"template_id": first["id"]}).status_code == 200
    assert run(client, task["id"], model) is not None
    assert len(model.calls) == 2


def test_large_pdf_capacity_and_actual_model_images_use_selected_pages(client, tmp_path):
    selected = template(client, complete=False)
    limited(client)
    pdf_path = tmp_path / "eighty.pdf"
    document = pdfium.PdfDocument.new()
    for _ in range(80):
        page = document.new_page(32, 32)
        page.close()
    document.save(pdf_path)
    document.close()
    upload = client.post("/api/v1/tasks", files={"file": ("eighty.pdf", pdf_path.read_bytes(), "application/pdf")}, data={"template_id": selected["id"]})
    assert upload.status_code == 201, upload.text
    task = upload.json()
    assert task["page_count"] == 80
    rejected = client.post("/api/v1/tasks", files={"file": ("another.pdf", pdf_path.read_bytes(), "application/pdf")})
    assert rejected.status_code == 429  # four selected visual units, five active-unit budget

    class ImageModel:
        model_name = "pdf-test"

        def extract_images(self, paths, prompt, result_type):
            assert [Path(path).name for path in paths] == ["page-1.png", "page-2.png", "page-3.png", "page-4.png"]
            assert "1–4" in prompt
            return result_type.model_validate({"header": {"title": "测试"}, "items": []})

    result = run(client, task["id"], ImageModel())
    assert result is not None
    assert result.input_scope.selected_units == 4
    assert result.input_scope.total_units == 80
