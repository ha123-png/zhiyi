import zipfile

from openpyxl import Workbook
import pytest

from document_pipeline_api.config import Settings
from document_pipeline_api.schemas.input_scope import SourceRange
from document_pipeline_api.services.file_formats import UnsupportedTextFileError
from document_pipeline_api.services.input_scope import (
    InputBudgetError, InputUnit, native_units, select_input, visual_units,
)


@pytest.fixture
def settings(tmp_path):
    return Settings(database_url="sqlite://", storage_dir=tmp_path)


def test_pdf_head_tail_reports_actual_and_missing_pages():
    plan = select_input(visual_units(80), text_budget=0, image_budget=10)
    assert [(r.start, r.end) for r in plan.scope.selected] == [(1, 5), (76, 80)]
    assert [(r.location.start, r.location.end) for r in plan.scope.omitted] == [(6, 75)]
    assert plan.scope.selected_units == 10
    assert plan.scope.total_units == 80
    assert plan.scope.coverage == "partial"
    assert "不能声称全文摘要或全部明细" in plan.scope.prompt_notice()


def test_small_input_keeps_every_unit_without_duplicating_overlap():
    plan = select_input(visual_units(3, frames=True), text_budget=0, image_budget=10)
    assert len(plan.units) == 3
    assert plan.scope.coverage == "complete"
    assert plan.scope.selected[0].kind == "frame"
    assert plan.scope.omitted == []


def test_long_text_preserves_tail_and_original_line_numbers(tmp_path, settings):
    path = tmp_path / "long.md"
    path.write_text("\n".join(["# 开头元信息"] + [f"正文 {i}" for i in range(2000)] + ["# 结尾重要金额 987654"]), encoding="utf-8")
    before = path.read_bytes()
    plan = select_input(native_units(path, "text/markdown", settings), text_budget=500, image_budget=0)
    assert "开头元信息" in plan.text
    assert "结尾重要金额 987654" in plan.text
    assert plan.scope.selected[-1].end == 2002
    assert plan.scope.text_characters == len(plan.text) <= 500
    assert path.read_bytes() == before
    assert {r.kind for r in plan.scope.selected} == {"line"}


def test_long_single_line_is_not_clipped_or_reported_as_complete():
    units = [InputUnit(SourceRange(kind="line", start=1, end=1), text="甲" * 2000)]
    with pytest.raises(InputBudgetError, match="完整内容单元"):
        select_input(units, text_budget=100, image_budget=0)


def test_oversized_middle_unit_is_explicitly_omitted():
    units = [InputUnit(SourceRange(kind="paragraph", start=i, end=i), text=text)
             for i, text in enumerate(["开头", "甲" * 2000, "结尾"], 1)]
    plan = select_input(units, text_budget=100, image_budget=0)
    assert [u.text for u in plan.units] == ["开头", "结尾"]
    assert plan.scope.omitted[0].location == SourceRange(kind="paragraph", start=2, end=2)


def test_xlsx_multiple_sheets_large_tail_and_formula(tmp_path, settings):
    path = tmp_path / "rows.xlsx"
    book = Workbook(write_only=True)
    first = book.create_sheet("前表")
    first.append(["名称", "金额"])
    for i in range(20000):
        first.append([f"记录 {i}", i])
    last = book.create_sheet("结尾表")
    last.append([None, None])
    last.append(["尾部凭证", "=SUM(1,2)"])
    book.save(path)
    plan = select_input(native_units(path, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", settings), text_budget=1000, image_budget=0)
    assert "前表!A1:B" in plan.text
    assert "结尾表!A2:B2" in plan.text
    assert "尾部凭证" in plan.text and "=SUM(1,2)" in plan.text
    assert plan.scope.total_units == 20003
    assert all(r.kind == "sheet_row" for r in plan.scope.selected)
    assert plan.scope.coverage == "partial"


def test_xlsx_unreliable_dimension_cannot_hide_late_rows(tmp_path, settings):
    path = tmp_path / "dimensions.xlsx"
    book = Workbook()
    book.active["A1"] = "开头"
    book.active["H50"] = "不能漏掉"
    book.save(path)
    with zipfile.ZipFile(path) as source:
        parts = {name: source.read(name) for name in source.namelist()}
    parts["xl/worksheets/sheet1.xml"] = parts["xl/worksheets/sheet1.xml"].replace(b'A1:H50', b'A1:A1')
    with zipfile.ZipFile(path, "w") as target:
        for name, data in parts.items():
            target.writestr(name, data)
    units = list(native_units(path, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", settings))
    assert len(units) == 50
    assert units[-1].location.columns == 8
    assert "不能漏掉" in units[-1].text


def test_docx_paragraphs_tables_and_images_are_distinct(tmp_path, settings):
    path = tmp_path / "mixed.docx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", '''<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><w:body>
<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>章节标题</w:t></w:r></w:p>
<w:tbl><w:tr><w:tc><w:p><w:r><w:t>表格内容</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
<w:p><w:r><w:t>最后一段</w:t><w:drawing><a:blip r:embed="rId1"/></w:drawing></w:r></w:p>
</w:body></w:document>''')
        archive.writestr("word/_rels/document.xml.rels", '<Relationships><Relationship Id="rId1" Target="media/image1.png"/></Relationships>')
        archive.writestr("word/media/image1.png", b"rendering validates actual pixels later")
    plan = select_input(native_units(path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", settings), text_budget=2000, image_budget=3)
    assert [u.location.kind for u in plan.units] == ["paragraph", "table_row", "paragraph"]
    assert "表格 1 · 第 1 行" in plan.text
    assert plan.scope.coverage == "partial"
    assert plan.scope.omitted[0].reason == "images_disabled"
    enabled = select_input(native_units(path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", settings, include_images=True), text_budget=2000, image_budget=3)
    assert enabled.units[-1].image_member == "word/media/image1.png"
    assert enabled.scope.coverage == "complete"


def test_invalid_text_tail_is_validated_before_selection(tmp_path, settings):
    path = tmp_path / "bad.txt"
    path.write_bytes(b"a\n" * 1000 + b"\xff")
    with pytest.raises(UnsupportedTextFileError, match="编码"):
        select_input(native_units(path, "text/plain", settings), text_budget=30, image_budget=0)


def test_zip_safety_budget_is_not_bypassed_by_small_input_budget(tmp_path, settings):
    path = tmp_path / "oversized.docx"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "x" * 1000)
    from dataclasses import replace

    with pytest.raises(UnsupportedTextFileError, match="解压后"):
        select_input(native_units(path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", replace(settings, max_import_uncompressed_bytes=100)), text_budget=30, image_budget=0)


@pytest.mark.parametrize("budget", [1, 2, 3, 5, 9, 10, 11, 20])
def test_visual_scope_partitions_original_without_gaps_or_overlap(budget):
    plan = select_input(visual_units(11), text_budget=0, image_budget=budget)
    selected = {n for r in plan.scope.selected for n in range(r.start, r.end + 1)}
    omitted = {n for r in plan.scope.omitted for n in range(r.location.start, r.location.end + 1)}
    assert selected.isdisjoint(omitted)
    assert selected | omitted == set(range(1, 12))
    assert len(selected) == min(budget, 11)


def test_empty_selected_boundaries_are_not_sent_as_useful_input():
    units = [InputUnit(SourceRange(kind="line", start=i, end=i), text=text)
             for i, text in enumerate(["", "重要内容" * 1000, ""], 1)]
    with pytest.raises(InputBudgetError):
        select_input(units, text_budget=100, image_budget=0)


def test_word_unhandled_parts_cannot_be_reported_as_full_coverage(tmp_path, settings):
    path = tmp_path / "header.docx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>正文</w:t><w:tab/><w:t>金额</w:t><w:br/><w:t>100</w:t></w:r></w:p><w:sdt><w:t>内容控件</w:t></w:sdt></w:body></w:document>')
        archive.writestr("word/header1.xml", "页眉")
    plan = select_input(native_units(path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", settings), text_budget=1000, image_budget=0)
    assert "正文\t金额\n100" in plan.text
    assert plan.scope.coverage == "partial"
    assert len(plan.scope.omitted) == 2
    assert {r.reason for r in plan.scope.omitted} == {"unsupported_structure"}
