from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from document_pipeline_api.config import Settings
from document_pipeline_api.main import create_app
from document_pipeline_api.models import DataRowRecord, DataTableRecord, TaskRecord
from document_pipeline_api.services.data_tables import table_columns


def _client(tmp_path: Path) -> TestClient:
    settings = Settings(
        database_url=f"sqlite:///{tmp_path / 'tables.db'}",
        storage_dir=tmp_path / "uploads",
        queue_enabled=False,
    )
    return TestClient(create_app(settings))


def _create_manual_table(client: TestClient, name: str = "我的发票表") -> dict:
    response = client.post(
        "/api/v1/tables",
        json={"name": name, "template_key": "builtin-invoice"},
    )
    assert response.status_code == 201, response.text
    return response.json()


def _add_row(client: TestClient, table_id: str, **values) -> dict:
    response = client.post(
        f"/api/v1/tables/{table_id}/rows",
        json={"values": values},
    )
    assert response.status_code == 201, response.text
    return response.json()


def test_manual_empty_row_edit_allows_missing_field(tmp_path) -> None:
    """回归：新增空行后，编辑表列存在但该行缺失的字段不应被拒绝
    （曾报"数据表中不存在字段"/"缺少来源任务，不能修改"）。"""
    with _client(tmp_path) as client:
        table = _create_manual_table(client)
        row = _add_row(client, table["id"])
        assert row["values"] == {}
        assert row["task_id"] is None

        response = client.patch(
            f"/api/v1/tables/{table['id']}/rows/{row['id']}",
            json={
                "expected_version": row["version"],
                "changes": {"seller_name": "甲公司", "total_amount": 372},
            },
        )
        assert response.status_code == 200, response.text
        updated = response.json()
        assert updated["values"]["seller_name"] == "甲公司"
        assert updated["values"]["total_amount"] == 372
        assert updated["version"] == row["version"] + 1

        revisions = client.get(
            f"/api/v1/tables/{table['id']}/rows/{row['id']}/revisions",
        )
        assert revisions.status_code == 200
        latest = revisions.json()[-1]
        assert latest["operation"] == "table_edit"
        assert latest["before"] == {}
        assert latest["after"]["seller_name"] == "甲公司"
        assert latest["after"]["total_amount"] == 372


def test_merged_group_header_edit_is_visible_from_every_detail_revision(tmp_path) -> None:
    with _client(tmp_path) as client:
        table = _create_manual_table(client, "合并表")
        with client.app.state.session_factory() as session:
            rows = [
                DataRowRecord(
                    table_id=table["id"],
                    task_id=None,
                    item_index=index,
                    row_json=(
                        '{"__row_group":"source:1","seller_name":"甲公司",'
                        f'"name":"货品{index}"}}'
                    ),
                    row_version=1,
                )
                for index in (1, 2)
            ]
            session.add_all(rows)
            session.commit()
            row_ids = [row.id for row in rows]

        response = client.patch(
            f"/api/v1/tables/{table['id']}/rows/{row_ids[0]}",
            json={
                "expected_version": 1,
                "changes": {"seller_name": "乙公司"},
            },
        )
        assert response.status_code == 200, response.text

        detail = client.get(f"/api/v1/tables/{table['id']}").json()
        assert [row["values"]["seller_name"] for row in detail["rows"]] == [
            "乙公司",
            "乙公司",
        ]
        for row_id in row_ids:
            revisions = client.get(
                f"/api/v1/tables/{table['id']}/rows/{row_id}/revisions"
            ).json()
            assert revisions[-1]["before"]["seller_name"] == "甲公司"
            assert revisions[-1]["after"]["seller_name"] == "乙公司"


def test_manual_table_create_columns_and_row_crud(tmp_path) -> None:
    with _client(tmp_path) as client:
        table = _create_manual_table(client)
        assert table["row_count"] == 0

        detail = client.get(f"/api/v1/tables/{table['id']}").json()
        labels = [column["label"] for column in detail["columns"]]
        assert "销售方" in labels
        assert "来源文件" not in labels
        assert "明细序号" not in labels

        row = _add_row(client, table["id"], seller_name="甲公司", total_amount=100)
        assert row["values"]["seller_name"] == "甲公司"
        assert row["task_id"] is None

        renamed = client.patch(
            f"/api/v1/tables/{table['id']}",
            json={"name": "新名字"},
        )
        assert renamed.status_code == 200
        assert renamed.json()["name"] == "新名字"

        hit = client.get(
            f"/api/v1/tables/{table['id']}",
            params={"search": "甲公司"},
        ).json()
        assert hit["row_count"] == 1
        miss = client.get(
            f"/api/v1/tables/{table['id']}",
            params={"search": "不存在的关键字"},
        ).json()
        assert miss["row_count"] == 0

        deleted = client.request(
            "DELETE",
            f"/api/v1/tables/{table['id']}/rows",
            json={"row_ids": [row["id"]]},
        )
        assert deleted.status_code == 204
        empty = client.get(f"/api/v1/tables/{table['id']}").json()
        assert empty["row_count"] == 0


def test_manual_row_rejects_nested_values(tmp_path) -> None:
    with _client(tmp_path) as client:
        table = _create_manual_table(client)
        response = client.post(
            f"/api/v1/tables/{table['id']}/rows",
            json={"values": {"seller_name": {"nested": True}}},
        )
        assert response.status_code == 422


def test_manual_row_drops_fields_outside_column_contract(tmp_path) -> None:
    """新增行只接受列契约内的字段：契约外键（如被裁剪的幽灵列）静默丢弃。"""
    with _client(tmp_path) as client:
        table = _create_manual_table(client)
        row = _add_row(
            client,
            table["id"],
            seller_name="甲公司",
            total_amount=100,
            ghost_column="不应存在",
        )
        assert row["values"].get("seller_name") == "甲公司"
        assert "ghost_column" not in row["values"]


def test_merge_tables_union_columns_and_rows(tmp_path) -> None:
    with _client(tmp_path) as client:
        first = _create_manual_table(client, "表一")
        _add_row(client, first["id"], seller_name="甲公司", total_amount=100)
        second = _create_manual_table(client, "表二")
        _add_row(client, second["id"], seller_name="乙公司", total_amount=200)

        response = client.post(
            "/api/v1/tables/merge",
            json={"table_ids": [first["id"], second["id"]], "name": "合并表"},
        )
        assert response.status_code == 201, response.text
        merged = response.json()
        assert merged["name"] == "合并表"
        assert merged["row_count"] == 2

        detail = client.get(f"/api/v1/tables/{merged['id']}").json()
        assert detail["row_count"] == 2
        seller_key = next(column["key"] for column in detail["columns"] if column["label"] == "销售方")
        sellers = {row["values"].get(seller_key) for row in detail["rows"]}
        assert sellers == {"甲公司", "乙公司"}


def test_merge_tables_unifies_same_visible_label_with_different_internal_keys(tmp_path) -> None:
    with _client(tmp_path) as client:
        first = client.post(
            "/api/v1/tables/import-new",
            files={"file": ("a.csv", "销售方\n甲公司\n".encode(), "text/csv")},
        ).json()
        second = client.post(
            "/api/v1/tables/import-new",
            files={"file": ("b.csv", "销售方\n乙公司\n".encode(), "text/csv")},
        ).json()
        merged = client.post(
            "/api/v1/tables/merge",
            json={"table_ids": [first["id"], second["id"]], "name": "同名字段合并"},
        )
        assert merged.status_code == 201, merged.text
        detail = client.get(f"/api/v1/tables/{merged.json()['id']}").json()

    assert [column["label"] for column in detail["columns"]] == ["销售方"]
    canonical_key = detail["columns"][0]["key"]
    assert {row["values"][canonical_key] for row in detail["rows"]} == {"甲公司", "乙公司"}


def test_merge_tables_is_column_union_plus_vertical_row_append(tmp_path) -> None:
    """合并不是 JOIN：同名列合一、所有来源行纵向追加、缺失列留空。"""
    with _client(tmp_path) as client:
        students = client.post(
            "/api/v1/tables/import-new",
            files={"file": ("students.csv", "学号,姓名\n2024,张三\n2025,李四\n".encode(), "text/csv")},
        ).json()
        genders = client.post(
            "/api/v1/tables/import-new",
            files={"file": ("genders.csv", "姓名,性别\n张三,男\n李四,男\n".encode(), "text/csv")},
        ).json()
        merged = client.post(
            "/api/v1/tables/merge",
            json={"table_ids": [students["id"], genders["id"]], "name": "学生合并"},
        )
        assert merged.status_code == 201, merged.text
        detail = client.get(f"/api/v1/tables/{merged.json()['id']}").json()

    assert [column["label"] for column in detail["columns"]] == ["学号", "姓名", "性别"]
    keys = {column["label"]: column["key"] for column in detail["columns"]}
    values = [row["values"] for row in detail["rows"]]
    assert len(values) == 4
    assert [(row.get(keys["学号"]), row.get(keys["姓名"]), row.get(keys["性别"])) for row in values] == [
        ("2024", "张三", None),
        ("2025", "李四", None),
        (None, "张三", "男"),
        (None, "李四", "男"),
    ]


def test_merge_tables_rejects_historical_duplicate_visible_labels(tmp_path) -> None:
    with _client(tmp_path) as client:
        first_response = client.post(
            "/api/v1/tables/import-new",
            files={"file": ("a.csv", "供货方,供货方,收货方\n甲公司,,乙公司\n".encode(), "text/csv")},
        )
    assert first_response.status_code == 422
    assert "重复列名" in first_response.json()["detail"]


def test_delete_table_cascades_rows_and_views(tmp_path) -> None:
    with _client(tmp_path) as client:
        table = _create_manual_table(client)
        _add_row(client, table["id"], seller_name="甲公司")
        _add_row(client, table["id"], seller_name="乙公司")
        split = client.post(
            f"/api/v1/tables/{table['id']}/split",
            json={"field_key": "seller_name"},
        )
        assert split.status_code == 200, split.text
        views = client.get(f"/api/v1/tables/{table['id']}/views").json()
        assert len(views) == 2

        deleted = client.delete(f"/api/v1/tables/{table['id']}")
        assert deleted.status_code == 204
        assert client.get(f"/api/v1/tables/{table['id']}").status_code == 404
        assert client.get(f"/api/v1/tables/{table['id']}/views").status_code == 404


def test_export_csv_and_json_use_chinese_headers(tmp_path) -> None:
    with _client(tmp_path) as client:
        table = _create_manual_table(client)
        _add_row(client, table["id"], seller_name="甲公司", total_amount=100)

        csv_response = client.get(f"/api/v1/tables/{table['id']}/export.csv")
        assert csv_response.status_code == 200
        content = csv_response.content.decode("utf-8-sig")
        assert "销售方" in content

        json_response = client.get(f"/api/v1/tables/{table['id']}/export.json")
        assert json_response.status_code == 200
        payload = json_response.json()
        assert payload[0]["seller_name"] == "甲公司"


def test_manual_table_columns_keep_template_contract_after_partial_row(tmp_path) -> None:
    """模板列是固定结构，不能因当前行没填值而消失。"""
    with _client(tmp_path) as client:
        table = _create_manual_table(client)
        detail = client.get(f"/api/v1/tables/{table['id']}").json()
        # 空表保留完整模板契约，供录入
        assert any(column["label"] == "开票日期" for column in detail["columns"])

        _add_row(client, table["id"], seller_name="上海A公司", total_amount=100.5)

        detail = client.get(f"/api/v1/tables/{table['id']}").json()
        keys = [column["key"] for column in detail["columns"]]
        assert "seller_name" in keys
        assert "total_amount" in keys
        assert "document_date" in keys
        assert "开票日期" in [column["label"] for column in detail["columns"]]


def test_stats_summary_returns_average_elapsed_and_failure_rate(tmp_path) -> None:
    from document_pipeline_api.models import ExtractionRecord, TaskRecord
    from document_pipeline_api.models.task import utc_now

    with _client(tmp_path) as client:
        with client.app.state.session_factory() as session:
            for index, task_id in enumerate(["t1", "t2"], start=1):
                session.add(
                    TaskRecord(
                        id=task_id,
                        filename=f"{index}.png",
                        content_type="image/png",
                        size_bytes=1,
                        page_count=1,
                        sha256=f"digest-{index}",
                        storage_path=f"{index}.png",
                        template_mode="invoice",
                        status="completed" if task_id == "t1" else "failed",
                        created_at=utc_now(),
                        updated_at=utc_now(),
                    )
                )
            session.flush()
            for index, task_id in enumerate(["t1", "t2"], start=1):
                session.add(
                    ExtractionRecord(
                        task_id=task_id,
                        document_kind="invoice",
                        model_name="qwen-test",
                        prompt_version="p1",
                        elapsed_seconds=10.0 * index,
                        result_json="{}",
                        validation_json="[]",
                        evidence_json="[]",
                        created_at=utc_now(),
                    )
                )
            session.commit()

        summary = client.get("/api/v1/stats/summary").json()
        assert summary["average_elapsed_seconds"] == 15.0
        assert summary["processed_count"] == 2
        assert summary["failed_count"] == 1
        assert summary["success_rate"] == 0.5


def test_builtin_table_columns_derive_from_template(tmp_path) -> None:
    """内置表列契约按模板字段生成：中文列名与模板一致，item 金额/税额
    别名到真实存储键（amount→item_amount、tax_amount→item_tax_amount）。"""
    from sqlalchemy.orm import sessionmaker

    from document_pipeline_api.db import build_engine
    from document_pipeline_api.migrations import upgrade_database
    from document_pipeline_api.services.templates import ensure_builtin_templates

    settings = Settings(
        database_url=f"sqlite:///{tmp_path / 'columns.db'}",
        storage_dir=tmp_path / "uploads",
        queue_enabled=False,
    )
    engine = build_engine(settings.database_url)
    upgrade_database(engine)
    session_factory = sessionmaker(engine, expire_on_commit=False)
    with session_factory() as session:
        ensure_builtin_templates(session)
        table = DataTableRecord(
            id="builtin-table",
            name="发票",
            template_key="invoice",
            template_version="builtin-v1",
            document_kind="invoice",
        )
        session.add(table)
        session.commit()
        columns = table_columns(session, table)
    labels = [column.label for column in columns]
    keys = [column.key for column in columns]
    assert "销售方" in labels
    assert "购买方" in labels
    assert "总税额" in labels
    assert labels.count("税额") == 1
    assert "来源文件" not in labels
    assert "明细序号" not in labels
    # 模板 item 键别名到真实存储键，行数据才能对上
    assert "item_amount" in keys
    assert "item_tax_amount" in keys
    # 模板字段（不含 document_type）生效
    assert "document_type" not in keys


def test_xlsx_export_still_works_for_manual_table(tmp_path) -> None:
    with _client(tmp_path) as client:
        table = _create_manual_table(client)
        _add_row(client, table["id"], seller_name="甲公司", total_amount=100)
        response = client.get(f"/api/v1/tables/{table['id']}/export.xlsx")
        assert response.status_code == 200
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        assert "销售方" in headers


def test_import_csv_appends_rows_by_chinese_headers(tmp_path) -> None:
    with _client(tmp_path) as client:
        table = _create_manual_table(client)
        content = "销售方,合计金额\n甲公司,100\n乙公司,200\n".encode("utf-8-sig")
        response = client.post(
            f"/api/v1/tables/{table['id']}/import",
            files={"file": ("import.csv", content, "text/csv")},
        )
        assert response.status_code == 200, response.text

        detail = client.get(f"/api/v1/tables/{table['id']}").json()
        assert detail["row_count"] == 2
        sellers = {row["values"].get("seller_name") for row in detail["rows"]}
        assert sellers == {"甲公司", "乙公司"}


def test_import_rejects_non_xlsx_csv(tmp_path) -> None:
    with _client(tmp_path) as client:
        table = _create_manual_table(client)
        response = client.post(
            f"/api/v1/tables/{table['id']}/import",
            files={"file": ("data.txt", b"hello", "text/plain")},
        )
        assert response.status_code == 422


def test_import_rejects_oversized_file_before_parsing(tmp_path) -> None:
    settings = Settings(
        database_url=f"sqlite:///{tmp_path / 'limited-import.db'}",
        storage_dir=tmp_path / "uploads",
        max_import_bytes=8,
        queue_enabled=False,
    )
    with TestClient(create_app(settings)) as client:
        table = _create_manual_table(client)
        response = client.post(
            f"/api/v1/tables/{table['id']}/import",
            files={"file": ("import.csv", b"123456789", "text/csv")},
        )

    assert response.status_code == 413
    assert "超过" in response.json()["detail"]


def test_import_rejects_bad_workbook_and_duplicate_headers(tmp_path) -> None:
    with _client(tmp_path) as client:
        table = _create_manual_table(client)
        broken = client.post(
            f"/api/v1/tables/{table['id']}/import",
            files={"file": ("broken.xlsx", b"not-a-workbook", "application/octet-stream")},
        )
        duplicate = client.post(
            f"/api/v1/tables/{table['id']}/import",
            files={"file": ("duplicate.csv", "销售方,销售方\n甲,乙\n".encode(), "text/csv")},
        )

    assert broken.status_code == 422
    assert "损坏" in broken.json()["detail"]
    assert duplicate.status_code == 422
    assert "重复列名" in duplicate.json()["detail"]


def test_import_rejects_duplicate_new_columns_instead_of_losing_values(tmp_path) -> None:
    with _client(tmp_path) as client:
        table = _create_manual_table(client)
        content = "销售方,新备注,新备注\n甲公司,加急,夜间配送\n".encode("utf-8-sig")
        response = client.post(
            f"/api/v1/tables/{table['id']}/import",
            files={"file": ("union.csv", content, "text/csv")},
        )
        assert response.status_code == 422, response.text
        assert "新备注" in response.json()["detail"]


def test_import_xlsx_expands_merged_cell_values_to_all_covered_rows(tmp_path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["销售方", "商品"])
    sheet.append(["甲公司", "商品一"])
    sheet.append([None, "商品二"])
    sheet.merge_cells("A2:A3")
    content = BytesIO()
    workbook.save(content)
    workbook.close()

    with _client(tmp_path) as client:
        table = _create_manual_table(client)
        response = client.post(
            f"/api/v1/tables/{table['id']}/import",
            files={
                "file": (
                    "merged.xlsx",
                    content.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert response.status_code == 200, response.text
        detail = client.get(f"/api/v1/tables/{table['id']}").json()

    assert [row["values"]["seller_name"] for row in detail["rows"]] == ["甲公司", "甲公司"]
    product_key = next(column["key"] for column in detail["columns"] if column["label"] == "商品")
    assert [row["values"][product_key] for row in detail["rows"]] == ["商品一", "商品二"]
    seller_column = next(column for column in detail["columns"] if column["label"] == "销售方")
    assert seller_column["section"] == "header"
    assert detail["rows"][0]["values"]["__row_group"].startswith("xlsx:")

    with _client(tmp_path) as client:
        # Use a new client only for a focused new-table round-trip below.
        response = client.post(
            "/api/v1/tables/import-new",
            files={
                "file": (
                    "merged.xlsx",
                    content.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        imported = response.json()
        exported = client.get(f"/api/v1/tables/{imported['id']}/export.xlsx")
        exported_book = load_workbook(BytesIO(exported.content))
        assert "A2:A3" in {
            str(cell_range) for cell_range in exported_book.active.merged_cells.ranges
        }
        exported_book.close()


def test_import_rejects_merged_or_incomplete_header(tmp_path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["组合表头", None])
    sheet.merge_cells("A1:B1")
    sheet.append(["甲", "乙"])
    content = BytesIO()
    workbook.save(content)
    workbook.close()
    with _client(tmp_path) as client:
        response = client.post(
            "/api/v1/tables/import-new",
            files={
                "file": (
                    "messy.xlsx",
                    content.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
    assert response.status_code == 422
    assert "单行表头" in response.json()["detail"]


def test_merge_real_task_backed_tables_does_not_reuse_unique_task_identity(tmp_path) -> None:
    with _client(tmp_path) as client:
        first = _create_manual_table(client)
        second = _create_manual_table(client, "第二表")
        with client.app.state.session_factory() as session:
            task = TaskRecord(
                id="source-task",
                filename="source.png",
                content_type="image/png",
                size_bytes=1,
                sha256="source-task",
                storage_path="source.png",
                template_mode="invoice",
                status="completed",
            )
            session.add(task)
            session.flush()
            session.add(
                DataRowRecord(
                    table_id=first["id"],
                    task_id="source-task",
                    item_index=1,
                    row_json='{"seller_name":"甲"}',
                    row_version=1,
                )
            )
            session.commit()
        _add_row(client, second["id"], seller_name="乙")
        merged = client.post(
            "/api/v1/tables/merge",
            json={"table_ids": [first["id"], second["id"]], "name": "合并表"},
        )
    assert merged.status_code == 201, merged.text
    detail = client.get(f"/api/v1/tables/{merged.json()['id']}").json()
    assert all(row["task_id"] is None for row in detail["rows"])


def test_import_enforces_row_limit(tmp_path) -> None:
    settings = Settings(
        database_url=f"sqlite:///{tmp_path / 'row-limit.db'}",
        storage_dir=tmp_path / "uploads",
        max_import_rows=1,
        queue_enabled=False,
    )
    with TestClient(create_app(settings)) as client:
        table = _create_manual_table(client)
        response = client.post(
            f"/api/v1/tables/{table['id']}/import",
            files={"file": ("import.csv", "销售方\n甲\n乙\n".encode(), "text/csv")},
        )

    assert response.status_code == 422
    assert "不能超过 1 行" in response.json()["detail"]
