from types import SimpleNamespace

from openpyxl import Workbook
import pytest

from document_pipeline_api.schemas.input_scope import InputScope, SourceRange
from document_pipeline_api.services.export_scope import WorkbookInputScopes, read_workbook_scopes


def test_scope_sheet_deduplicates_sources_and_preserves_all_long_notes():
    scope = InputScope(rule="all", coverage="complete", selected=[SourceRange(kind="line", start=1, end=2)], omitted=[], selected_units=2, total_units=2, text_characters=20, text_budget=100, image_budget=1, notes=["长" * 40000])
    record = SimpleNamespace(input_scope_json=scope.model_dump_json(), review_pending=None, row_json='{"source_filename":"=external.xlsx"}', task_id="same-source", id=1)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "知意输入范围"
    scopes = WorkbookInputScopes()
    for row in (2, 3):
        scopes.record(sheet, row, record, ["header", "item"], {"header"})
    scopes.finish(workbook)
    assert workbook.sheetnames == ["知意输入范围", "知意输入范围-2"]
    assert sheet.cell(2, 1).comment is None
    assert "来源编号 1" in sheet.cell(2, 2).comment.text
    assert "来源编号 1" in sheet.cell(3, 2).comment.text
    entries = list(workbook.worksheets[1].values)
    assert all(row[1] == "'=external.xlsx" for row in entries[2:] if row[3] != "内部范围记录 v1")
    assert sum(row[3] == "规则" for row in entries) == 1
    assert "".join(row[4] for row in entries if row[3] in {"说明", "说明（续）"}) == "长" * 40000
    restored = read_workbook_scopes(workbook, sheet)
    assert InputScope.model_validate_json(restored[2]) == scope
    assert restored[2] == restored[3]
    internal_rows = [index for index, row in enumerate(entries, 1) if row[3] == "内部范围记录 v1"]
    assert all(workbook.worksheets[1].row_dimensions[index].hidden for index in internal_rows)
    workbook.worksheets[1].delete_rows(internal_rows[0])
    with pytest.raises(ValueError, match="缺少片段"):
        read_workbook_scopes(workbook, sheet)


def test_legacy_rows_do_not_gain_a_false_complete_label():
    workbook = Workbook()
    scopes = WorkbookInputScopes()
    scopes.record(workbook.active, 2, SimpleNamespace(input_scope_json=None, review_pending=None), ["value"], set())
    scopes.finish(workbook)
    assert len(workbook.sheetnames) == 1
    assert workbook.active.cell(2, 1).comment is None
