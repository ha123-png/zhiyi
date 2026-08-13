from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from document_pipeline_api.config import Settings
from document_pipeline_api.main import create_app
from document_pipeline_api.models import ExtractionRecord, TaskRecord
from document_pipeline_api.schemas.extraction import DocumentExtraction, TemplateExtraction
from document_pipeline_api.services.data_tables import confirm_task
from document_pipeline_api.services.task_leases import TaskLeaseLostError


def _result(*, item_name: str = "服务") -> DocumentExtraction:
    return DocumentExtraction(
        document_type="发票",
        seller_name="销售方",
        buyer_name="购买方",
        document_number="NO-1",
        document_date="2026-07-30",
        amount_before_tax=100,
        tax_amount=6,
        total_amount=106,
        items=[
            {
                "name": item_name,
                "specification": "标准",
                "unit": "项",
                "quantity": 1,
                "unit_price": 100,
                "amount": 100,
                "tax_rate": "6%",
                "tax_amount": 6,
            }
        ],
    )


def _two_item_result() -> DocumentExtraction:
    result = _result(item_name="服务 A")
    result.items[0].quantity = 1
    result.items[0].unit_price = 50
    result.items[0].amount = 50
    result.items.append(
        result.items[0].model_copy(
            update={"name": "服务 B", "amount": 50}
        )
    )
    return result


def _seed_task(
    client: TestClient,
    task_id: str,
    *,
    item_name: str = "服务",
) -> None:
    with client.app.state.session_factory() as session:
        session.add(
            TaskRecord(
                id=task_id,
                filename=f"{task_id}.png",
                content_type="image/png",
                size_bytes=1,
                sha256=f"digest-{task_id}",
                storage_path=f"{task_id}.png",
                template_mode="invoice",
                status="needs_review",
            )
        )
        session.flush()
        session.add(
            ExtractionRecord(
                task_id=task_id,
                document_kind="invoice",
                model_name="test-model",
                prompt_version="test-prompt",
                elapsed_seconds=1,
                result_json=_result(item_name=item_name).model_dump_json(),
                validation_json="[]",
            )
        )
        session.commit()


def _seed_two_item_task(client: TestClient, task_id: str) -> None:
    with client.app.state.session_factory() as session:
        session.add(
            TaskRecord(
                id=task_id,
                filename=f"{task_id}.png",
                content_type="image/png",
                size_bytes=1,
                sha256=f"digest-{task_id}",
                storage_path=f"{task_id}.png",
                template_mode="invoice",
                status="needs_review",
            )
        )
        session.flush()
        session.add(
            ExtractionRecord(
                task_id=task_id,
                document_kind="invoice",
                model_name="test-model",
                prompt_version="test-prompt",
                elapsed_seconds=1,
                result_json=_two_item_result().model_dump_json(),
                validation_json="[]",
            )
        )
        session.commit()


def test_confirm_with_empty_items_inserts_header_only_row(tmp_path: Path) -> None:
    """模型对无明细文档返回 items=[] 时不应导致确认失败（历史 bug：item_index=None 被拒）。"""
    settings = Settings(
        database_url=f"sqlite:///{tmp_path / 'confirm-empty.db'}",
        storage_dir=tmp_path / "uploads",
    )
    with TestClient(create_app(settings)) as client:
        _seed_task(client, "empty-items")
        with client.app.state.session_factory() as session:
            extraction = session.get(ExtractionRecord, "empty-items")
            result = DocumentExtraction.model_validate_json(extraction.result_json)
            result.items = []
            extraction.result_json = result.model_dump_json()
            session.commit()

        response = client.post(
            "/api/v1/tasks/empty-items/confirm",
            json={"expected_review_version": 0},
        )

        assert response.status_code == 200
        confirmation = response.json()
        # 只有抬头、无明细 → 落 1 行 header-only 记录（item_index=1），不抛 RuntimeError
        assert confirmation["row_count"] == 1
        table = client.get(f"/api/v1/tables/{confirmation['table_id']}").json()
        assert table["row_count"] == 1
        assert table["rows"][0]["values"]["seller_name"] == "销售方"


def test_confirm_inserts_rows_and_exports_excel(tmp_path: Path) -> None:
    settings = Settings(
        database_url=f"sqlite:///{tmp_path / 'confirm.db'}",
        storage_dir=tmp_path / "uploads",
    )
    with TestClient(create_app(settings)) as client:
        _seed_task(client, "first", item_name="=unsafe")

        response = client.post(
            "/api/v1/tasks/first/confirm",
            json={"expected_review_version": 0},
        )

        assert response.status_code == 200
        confirmation = response.json()
        assert confirmation["row_count"] == 1
        assert client.get("/api/v1/tasks").json()[0]["status"] == "completed"

        table = client.get(f"/api/v1/tables/{confirmation['table_id']}").json()
        assert table["row_count"] == 1
        assert table["rows"][0]["values"]["seller_name"] == "销售方"
        assert table["rows"][0]["values"]["name"] == "=unsafe"

        export = client.get(
            f"/api/v1/tables/{confirmation['table_id']}/export.xlsx"
        )
        assert export.status_code == 200
        workbook = load_workbook(BytesIO(export.content))
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        # 内置表列契约与最终确认的原型字段逐字一致。
        assert headers[:4] == ["销售方", "购买方", "发票号码", "开票日期"]
        assert headers.count("总税额") == 1
        assert headers.count("税额") == 1
        assert sheet.cell(row=2, column=headers.index("商品名称") + 1).value == "'=unsafe"


def test_custom_columns_are_table_only_and_header_column_merges_on_export(tmp_path: Path) -> None:
    settings = Settings(
        database_url=f"sqlite:///{tmp_path / 'custom-columns.db'}",
        storage_dir=tmp_path / "uploads",
    )
    with TestClient(create_app(settings)) as client:
        _seed_task(client, "custom-columns", item_name="第一项")
        with client.app.state.session_factory() as session:
            extraction = session.get(ExtractionRecord, "custom-columns")
            assert extraction is not None
            extraction.result_json = _result(item_name="第一项").model_copy(
                update={"items": [
                    _result(item_name="第一项").items[0],
                    _result(item_name="第二项").items[0],
                ]}
            ).model_dump_json()
            session.commit()
        confirmation = client.post(
            "/api/v1/tasks/custom-columns/confirm",
            json={"expected_review_version": 0},
        ).json()
        table_id = confirmation["table_id"]

        added = client.post(
            f"/api/v1/tables/{table_id}/columns",
            json={"label": "内部备注", "section": "header"},
        )
        assert added.status_code == 201
        column = added.json()
        assert column["user_defined"] is True
        detail = client.get(f"/api/v1/tables/{table_id}").json()
        assert detail["columns"][-1]["label"] == "内部备注"

        edited = client.patch(
            f"/api/v1/tables/{table_id}/rows/{detail['rows'][0]['id']}",
            json={
                "expected_version": detail["rows"][0]["version"],
                "changes": {column["key"]: "待复核"},
            },
        )
        assert edited.status_code == 200
        refreshed = client.get(f"/api/v1/tables/{table_id}").json()
        assert all(row["values"][column["key"]] == "待复核" for row in refreshed["rows"])

        workbook = load_workbook(BytesIO(client.get(
            f"/api/v1/tables/{table_id}/export.xlsx"
        ).content))
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        custom_index = headers.index("内部备注") + 1
        seller_index = headers.index("销售方") + 1
        assert str(sheet.merged_cells.ranges).find(
            f"{sheet.cell(2, custom_index).coordinate}:{sheet.cell(3, custom_index).coordinate}"
        ) >= 0
        assert str(sheet.merged_cells.ranges).find(
            f"{sheet.cell(2, seller_index).coordinate}:{sheet.cell(3, seller_index).coordinate}"
        ) >= 0
        assert sheet.cell(1, custom_index).fill.fgColor.rgb == "00F0FDF4"
        assert sheet.cell(1, custom_index).font.color.rgb == "00065F46"
        assert sheet.column_dimensions[sheet.cell(1, custom_index).column_letter].width >= 10

        assert client.delete(
            f"/api/v1/tables/{table_id}/columns/{column['key']}"
        ).status_code == 204
        assert all(item["label"] != "内部备注" for item in client.get(
            f"/api/v1/tables/{table_id}"
        ).json()["columns"])

        template_key = refreshed["template_key"]
        assert client.delete(
            f"/api/v1/tables/{table_id}/columns/seller_name"
        ).status_code == 409
        assert template_key == "invoice"


def test_manual_confirm_cannot_bypass_active_worker_lease(tmp_path: Path) -> None:
    settings = Settings(
        database_url=f"sqlite:///{tmp_path / 'active-worker.db'}",
        storage_dir=tmp_path / "uploads",
    )
    with TestClient(create_app(settings)) as client:
        _seed_task(client, "active-worker")
        with client.app.state.session_factory() as session:
            task = session.get(TaskRecord, "active-worker")
            assert task is not None
            task.status = "validating"
            task.lease_token = "current-worker-token"
            session.commit()

        response = client.post(
            "/api/v1/tasks/active-worker/confirm",
            json={"expected_review_version": 0},
        )

        assert response.status_code == 409
        assert response.json()["detail"] == "只有待确认或已完成任务可以手动保存到表。"
        with client.app.state.session_factory() as session:
            task = session.get(TaskRecord, "active-worker")
            assert task is not None
            assert task.status == "validating"
            assert task.lease_token == "current-worker-token"


def test_late_worker_rolls_back_pending_writes_when_task_is_already_confirmed(
    tmp_path: Path,
) -> None:
    settings = Settings(
        database_url=f"sqlite:///{tmp_path / 'already-confirmed.db'}",
        storage_dir=tmp_path / "uploads",
    )
    with TestClient(create_app(settings)) as client:
        _seed_task(client, "already-confirmed", item_name="新 Worker")
        confirmed = client.post(
            "/api/v1/tasks/already-confirmed/confirm",
            json={"expected_review_version": 0},
        )
        assert confirmed.status_code == 200

        with client.app.state.session_factory() as session:
            extraction = session.get(ExtractionRecord, "already-confirmed")
            assert extraction is not None
            extraction.result_json = _result(item_name="迟到旧 Worker").model_dump_json()
            with pytest.raises(TaskLeaseLostError):
                confirm_task(
                    session,
                    "already-confirmed",
                    0,
                    expected_lease_token="expired-token",
                )

        with client.app.state.session_factory() as session:
            extraction = session.get(ExtractionRecord, "already-confirmed")
            assert extraction is not None
            persisted = DocumentExtraction.model_validate_json(extraction.result_json)
            assert persisted.items[0].name == "新 Worker"


def test_same_template_appends_and_confirmation_is_idempotent(
    tmp_path: Path,
) -> None:
    settings = Settings(
        database_url=f"sqlite:///{tmp_path / 'append.db'}",
        storage_dir=tmp_path / "uploads",
    )
    with TestClient(create_app(settings)) as client:
        _seed_task(client, "first")
        _seed_task(client, "second")

        first = client.post(
            "/api/v1/tasks/first/confirm",
            json={"expected_review_version": 0},
        ).json()
        second = client.post(
            "/api/v1/tasks/second/confirm",
            json={"expected_review_version": 0},
        ).json()
        repeated = client.post(
            "/api/v1/tasks/first/confirm",
            json={"expected_review_version": 0},
        ).json()

        assert first["table_id"] == second["table_id"] == repeated["table_id"]
        assert client.get("/api/v1/tables").json()[0]["row_count"] == 2
        first_page = client.get(
            f"/api/v1/tables/{first['table_id']}?page=1&page_size=1"
        ).json()
        second_page = client.get(
            f"/api/v1/tables/{first['table_id']}?page=2&page_size=1"
        ).json()
        assert first_page["row_count"] == second_page["row_count"] == 2
        assert first_page["page"] == 1
        assert second_page["page"] == 2
        assert first_page["rows"][0]["id"] != second_page["rows"][0]["id"]


def test_editing_completed_result_updates_existing_table_rows(
    tmp_path: Path,
) -> None:
    settings = Settings(
        database_url=f"sqlite:///{tmp_path / 'edit-completed.db'}",
        storage_dir=tmp_path / "uploads",
    )
    with TestClient(create_app(settings)) as client:
        _seed_task(client, "editable")
        confirmation = client.post(
            "/api/v1/tasks/editable/confirm",
            json={"expected_review_version": 0},
        ).json()
        edited = _result(item_name="人工修正")

        response = client.put(
            "/api/v1/tasks/editable/review",
            json={
                "expected_version": 0,
                "result": edited.model_dump(mode="json"),
            },
        )

        assert response.status_code == 200
        table = client.get(
            f"/api/v1/tables/{confirmation['table_id']}"
        ).json()
        assert table["row_count"] == 1
        assert table["rows"][0]["values"]["name"] == "人工修正"
        refreshed = client.get(
            "/api/v1/tasks/editable/confirmation"
        ).json()
        assert refreshed["review_version"] == 1


def test_valid_review_updates_rows_and_completes_without_extra_confirmation(
    tmp_path: Path,
) -> None:
    settings = Settings(
        database_url=f"sqlite:///{tmp_path / 'stale-confirm.db'}",
        storage_dir=tmp_path / "uploads",
    )
    with TestClient(create_app(settings)) as client:
        _seed_task(client, "stale")
        review = _result(item_name="人工修正")
        assert client.put(
            "/api/v1/tasks/stale/review",
            json={
                "expected_version": 0,
                "result": review.model_dump(mode="json"),
            },
        ).status_code == 200

        response = client.post(
            "/api/v1/tasks/stale/confirm",
            json={"expected_review_version": 0},
        )

        assert response.status_code == 200
        table = client.get("/api/v1/tables").json()[0]
        detail = client.get(f"/api/v1/tables/{table['id']}").json()
        assert detail["rows"][0]["values"]["name"] == "人工修正"
        assert client.get("/api/v1/tasks").json()[0]["status"] == "completed"


def test_custom_template_flattens_fields_and_uses_output_mapping(
    tmp_path: Path,
) -> None:
    settings = Settings(
        database_url=f"sqlite:///{tmp_path / 'custom-table.db'}",
        storage_dir=tmp_path / "uploads",
    )
    with TestClient(create_app(settings)) as client:
        template = client.post(
            "/api/v1/templates",
            json={
                "name": "门店评分",
                "description": "整理门店评分",
                "extra_instructions": "",
                "fields": [
                    {
                        "key": "store",
                        "label": "门店",
                        "section": "header",
                    },
                    {
                        "key": "score",
                        "label": "得分",
                        "section": "item",
                        "value_type": "number",
                    },
                ],
                "validation_rules": [],
                "output_mapping": {"score": "评分"},
            },
        ).json()
        with client.app.state.session_factory() as session:
            session.add(
                TaskRecord(
                    id="custom-result",
                    filename="score.png",
                    content_type="image/png",
                    size_bytes=1,
                    sha256="custom-digest",
                    storage_path="score.png",
                    template_mode="smart",
                    template_id=template["id"],
                    template_version=1,
                    status="needs_review",
                )
            )
            session.flush()
            session.add(
                ExtractionRecord(
                    task_id="custom-result",
                    document_kind="custom",
                    template_id=template["id"],
                    template_version=1,
                    model_name="test-model",
                    prompt_version="test-prompt",
                    elapsed_seconds=1,
                    result_json=(
                        '{"header":{"store":"一号门店"},'
                        '"items":[{"score":98.5},{"score":96}]}'
                    ),
                    validation_json="[]",
                )
            )
            session.commit()

        confirmation = client.post(
            "/api/v1/tasks/custom-result/confirm",
            json={"expected_review_version": 0},
        ).json()

        assert confirmation["row_count"] == 2
        table = client.get(f"/api/v1/tables/{confirmation['table_id']}").json()
        assert table["template_key"] == template["id"]
        assert table["template_version"] == "1"
        assert table["rows"][1]["values"]["store"] == "一号门店"
        assert table["rows"][1]["values"]["score"] == 96

        export = client.get(
            f"/api/v1/tables/{confirmation['table_id']}/export.xlsx"
        )
        workbook = load_workbook(BytesIO(export.content))
        headers = [cell.value for cell in workbook.active[1]]
        assert headers == ["门店", "评分"]


def test_review_sync_keeps_row_identity_and_records_revision(tmp_path: Path) -> None:
    settings = Settings(
        database_url=f"sqlite:///{tmp_path / 'stable-row.db'}",
        storage_dir=tmp_path / "uploads",
    )
    with TestClient(create_app(settings)) as client:
        _seed_task(client, "stable")
        confirmation = client.post(
            "/api/v1/tasks/stable/confirm",
            json={"expected_review_version": 0},
        ).json()
        original = client.get(
            f"/api/v1/tables/{confirmation['table_id']}"
        ).json()["rows"][0]

        edited = _result(item_name="审核修正")
        response = client.put(
            "/api/v1/tasks/stable/review",
            json={
                "expected_version": 0,
                "result": edited.model_dump(mode="json"),
            },
        )

        assert response.status_code == 200
        current = client.get(
            f"/api/v1/tables/{confirmation['table_id']}"
        ).json()["rows"][0]
        assert current["id"] == original["id"]
        assert current["version"] == original["version"] + 1
        revisions = client.get(
            f"/api/v1/tables/{confirmation['table_id']}"
            f"/rows/{original['id']}/revisions"
        ).json()
        assert [revision["operation"] for revision in revisions] == [
            "materialize_create",
            "materialize_update",
        ]
        assert revisions[-1]["before"]["name"] == "服务"
        assert revisions[-1]["after"]["name"] == "审核修正"


def test_row_edit_uses_optimistic_concurrency(tmp_path: Path) -> None:
    settings = Settings(
        database_url=f"sqlite:///{tmp_path / 'row-conflict.db'}",
        storage_dir=tmp_path / "uploads",
    )
    with TestClient(create_app(settings)) as client:
        _seed_task(client, "conflict")
        table_id = client.post(
            "/api/v1/tasks/conflict/confirm",
            json={"expected_review_version": 0},
        ).json()["table_id"]
        row = client.get(f"/api/v1/tables/{table_id}").json()["rows"][0]

        first = client.patch(
            f"/api/v1/tables/{table_id}/rows/{row['id']}",
            json={
                "expected_version": row["version"],
                "changes": {"name": "第一次修改"},
            },
        )
        stale = client.patch(
            f"/api/v1/tables/{table_id}/rows/{row['id']}",
            json={
                "expected_version": row["version"],
                "changes": {"name": "旧页面覆盖"},
            },
        )

        assert first.status_code == 200
        assert first.json()["version"] == row["version"] + 1
        assert stale.status_code == 409
        current = client.get(f"/api/v1/tables/{table_id}").json()["rows"][0]
        assert current["values"]["name"] == "第一次修改"


def test_completed_row_edit_rejects_broken_builtin_rule(tmp_path: Path) -> None:
    settings = Settings(
        database_url=f"sqlite:///{tmp_path / 'row-rule.db'}",
        storage_dir=tmp_path / "uploads",
    )
    with TestClient(create_app(settings)) as client:
        _seed_task(client, "row-rule")
        table_id = client.post(
            "/api/v1/tasks/row-rule/confirm",
            json={"expected_review_version": 0},
        ).json()["table_id"]
        row = client.get(f"/api/v1/tables/{table_id}").json()["rows"][0]

        response = client.patch(
            f"/api/v1/tables/{table_id}/rows/{row['id']}",
            json={
                "expected_version": row["version"],
                "changes": {"item_amount": 999},
            },
        )

        assert response.status_code == 422
        assert "第 1 行：数量 1.0 × 单价 100.0 不等于金额 999.0" in response.json()[
            "detail"
        ]
        current = client.get(f"/api/v1/tables/{table_id}").json()["rows"][0]
        assert current["values"]["item_amount"] == 100
        assert current["version"] == row["version"]


def test_header_edit_propagates_to_all_rows_of_same_document(tmp_path: Path) -> None:
    settings = Settings(
        database_url=f"sqlite:///{tmp_path / 'header-propagation.db'}",
        storage_dir=tmp_path / "uploads",
    )
    with TestClient(create_app(settings)) as client:
        _seed_two_item_task(client, "header-propagation")
        table_id = client.post(
            "/api/v1/tasks/header-propagation/confirm",
            json={"expected_review_version": 0},
        ).json()["table_id"]
        rows = client.get(f"/api/v1/tables/{table_id}").json()["rows"]

        response = client.patch(
            f"/api/v1/tables/{table_id}/rows/{rows[1]['id']}",
            json={
                "expected_version": rows[1]["version"],
                "changes": {"seller_name": "统一后的销售方"},
            },
        )

        assert response.status_code == 200
        current = client.get(f"/api/v1/tables/{table_id}").json()["rows"]
        assert [row["values"]["seller_name"] for row in current] == [
            "统一后的销售方",
            "统一后的销售方",
        ]
        assert [row["version"] for row in current] == [2, 2]


def test_custom_template_rule_also_guards_fact_table_edits(tmp_path: Path) -> None:
    settings = Settings(
        database_url=f"sqlite:///{tmp_path / 'custom-row-rule.db'}",
        storage_dir=tmp_path / "uploads",
    )
    with TestClient(create_app(settings)) as client:
        template = client.post(
            "/api/v1/templates",
            json={
                "name": "门店评分",
                "description": "",
                "extra_instructions": "",
                "fields": [
                    {"key": "store", "label": "门店", "section": "header"},
                    {
                        "key": "score",
                        "label": "分数",
                        "section": "item",
                        "value_type": "number",
                    },
                ],
                "validation_rules": [],
                "deterministic_rules": [
                    {
                        "kind": "range",
                        "field": "items[].score",
                        "minimum": 0,
                        "maximum": 100,
                    }
                ],
                "output_mapping": {},
            },
        ).json()
        with client.app.state.session_factory() as session:
            session.add(
                TaskRecord(
                    id="custom-row-rule",
                    filename="score.png",
                    content_type="image/png",
                    size_bytes=1,
                    sha256="custom-row-rule-digest",
                    storage_path="score.png",
                    template_mode=template["id"],
                    template_id=template["id"],
                    template_version=template["version"],
                    status="needs_review",
                )
            )
            session.flush()
            session.add(
                ExtractionRecord(
                    task_id="custom-row-rule",
                    document_kind="custom",
                    template_id=template["id"],
                    template_version=template["version"],
                    model_name="test-model",
                    prompt_version="test-prompt",
                    elapsed_seconds=1,
                    result_json=TemplateExtraction(
                        header={"store": "一号门店"},
                        items=[{"score": None}],
                    ).model_dump_json(),
                    validation_json="[]",
                )
            )
            session.commit()
        table_id = client.post(
            "/api/v1/tasks/custom-row-rule/confirm",
            json={"expected_review_version": 0},
        ).json()["table_id"]
        row = client.get(f"/api/v1/tables/{table_id}").json()["rows"][0]

        response = client.patch(
            f"/api/v1/tables/{table_id}/rows/{row['id']}",
            json={
                "expected_version": row["version"],
                "changes": {"score": 120},
            },
        )

        assert response.status_code == 422
        assert "不能大于 100" in response.json()["detail"]
        wrong_type = client.patch(
            f"/api/v1/tables/{table_id}/rows/{row['id']}",
            json={
                "expected_version": row["version"],
                "changes": {"score": "不是数字"},
            },
        )
        assert wrong_type.status_code == 422
        assert "值类型不正确" in wrong_type.json()["detail"]


def test_split_view_reuses_source_rows_and_edits_same_fact(tmp_path: Path) -> None:
    settings = Settings(
        database_url=f"sqlite:///{tmp_path / 'split-view.db'}",
        storage_dir=tmp_path / "uploads",
    )
    with TestClient(create_app(settings)) as client:
        _seed_task(client, "company-a")
        _seed_task(client, "company-b")
        first = client.post(
            "/api/v1/tasks/company-a/confirm",
            json={"expected_review_version": 0},
        ).json()
        client.post(
            "/api/v1/tasks/company-b/confirm",
            json={"expected_review_version": 0},
        )
        views = client.post(
            f"/api/v1/tables/{first['table_id']}/split",
            json={"field_key": "source_filename"},
        ).json()

        assert len(views) == 2
        selected = next(
            view for view in views if view["field_value"] == "company-a.png"
        )
        detail = client.get(
            f"/api/v1/tables/{first['table_id']}/views/{selected['id']}"
        ).json()
        assert detail["row_count"] == 1
        view_row = detail["rows"][0]

        edited = client.patch(
            f"/api/v1/tables/{first['table_id']}/rows/{view_row['id']}",
            json={
                "expected_version": view_row["version"],
                "changes": {"name": "分 Sheet 修改"},
            },
        )

        assert edited.status_code == 200
        source_rows = client.get(
            f"/api/v1/tables/{first['table_id']}"
        ).json()["rows"]
        same_row = next(row for row in source_rows if row["id"] == view_row["id"])
        assert same_row["values"]["name"] == "分 Sheet 修改"
        refreshed_view = client.get(
            f"/api/v1/tables/{first['table_id']}/views/{selected['id']}"
        ).json()
        assert refreshed_view["rows"][0]["id"] == view_row["id"]
        assert refreshed_view["rows"][0]["values"]["name"] == "分 Sheet 修改"


def test_review_sync_does_not_overwrite_a_user_edited_table_row(
    tmp_path: Path,
) -> None:
    settings = Settings(
        database_url=f"sqlite:///{tmp_path / 'preserve-user-edit.db'}",
        storage_dir=tmp_path / "uploads",
    )
    with TestClient(create_app(settings)) as client:
        _seed_task(client, "preserved")
        table_id = client.post(
            "/api/v1/tasks/preserved/confirm",
            json={"expected_review_version": 0},
        ).json()["table_id"]
        row = client.get(f"/api/v1/tables/{table_id}").json()["rows"][0]
        assert client.patch(
            f"/api/v1/tables/{table_id}/rows/{row['id']}",
            json={
                "expected_version": row["version"],
                "changes": {"name": "表格中的最终值"},
            },
        ).status_code == 200

        review_result = _result(item_name="较晚保存的审核值")
        assert client.put(
            "/api/v1/tasks/preserved/review",
            json={
                "expected_version": 0,
                "result": review_result.model_dump(mode="json"),
            },
        ).status_code == 200

        current = client.get(f"/api/v1/tables/{table_id}").json()["rows"][0]
        assert current["id"] == row["id"]
        assert current["values"]["name"] == "表格中的最终值"
        assert current["version"] == row["version"] + 1
