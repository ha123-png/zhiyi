import json
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from document_pipeline_api.config import Settings
from document_pipeline_api.main import create_app
from document_pipeline_api.services.data_tables import create_table_from_import, parse_import_rows
from document_pipeline_api.services.extraction import process_task
from document_pipeline_api.services.export_scope import REVIEW_KEY, REVIEW_LABEL


class MissingOwner:
    model_name = "offline-review-fixture"

    def complete_text(self, prompt, result_type):
        return result_type.model_validate({"header": {"owner": None}, "items": []})


@pytest.mark.parametrize("suffix", ["xlsx", "csv", "json"])
@pytest.mark.parametrize("resolve_review", [True, False])
def test_pending_review_survives_export_import_merge_history_and_clears_on_review(tmp_path, suffix, resolve_review):
    settings = Settings(database_url=f"sqlite:///{tmp_path / 'test.db'}", storage_dir=tmp_path / "uploads", queue_enabled=False)
    with TestClient(create_app(settings)) as client:
        created = client.post("/api/v1/templates", json={"name": "需要负责人", "fields": [{"key": "owner", "label": "负责人", "section": "header"}], "deterministic_rules": [{"kind": "required", "field": "header.owner", "severity": "error"}]})
        assert created.status_code == 201, created.text
        template = created.json()
        original = "负责人未填写。".encode()
        task = client.post("/api/v1/tasks", files={"file": ("待核对.txt", original, "text/plain")}, data={"template_id": template["id"]}).json()
        with client.app.state.session_factory() as session:
            result = process_task(session, settings, task["id"], client=MissingOwner())
        assert result is not None and result.validation_issues
        table = client.get("/api/v1/tables").json()[0]
        table_id = table["id"]
        row = client.get(f"/api/v1/tables/{table_id}").json()["rows"][0]
        assert row["review_pending"] is True
        exported = client.get(f"/api/v1/tables/{table_id}/export.{suffix}")
        assert exported.status_code == 200
        if suffix == "json":
            imported_rows = exported.json()
            assert imported_rows[0][REVIEW_KEY] is True
        else:
            imported_rows = parse_import_rows(exported.content, f"rows.{suffix}")
            assert imported_rows[0][REVIEW_KEY] == "待核对"
            if suffix == "xlsx":
                workbook = load_workbook(BytesIO(exported.content))
                assert REVIEW_LABEL in [cell.value for cell in workbook.active[1]]
                assert "重新执行校验" in workbook.active.cell(2, 2).comment.text
                workbook.close()
        with client.app.state.session_factory() as session:
            imported = create_table_from_import(session, "重新导入", imported_rows)
            imported_id = imported.id
        imported_detail = client.get(f"/api/v1/tables/{imported_id}").json()
        assert imported_detail["rows"][0]["review_pending"] is True
        assert REVIEW_LABEL not in [col["label"] for col in imported_detail["columns"]]
        merged = client.post("/api/v1/tables/merge", json={"name": "合并", "table_ids": [table_id, imported_id]})
        assert merged.status_code == 201
        assert all(row["review_pending"] for row in client.get(f"/api/v1/tables/{merged.json()['id']}").json()["rows"])
        if resolve_review:
            updated = client.put(f"/api/v1/tasks/{task['id']}/review", json={"expected_version": 0, "result": {"header": {"owner": "合成负责人"}, "items": []}})
            assert updated.status_code == 200, updated.text
            assert client.get(f"/api/v1/tables/{table_id}").json()["rows"][0]["review_pending"] is False
            # An independent imported copy must keep the status it actually carried.
            assert client.get(f"/api/v1/tables/{imported_id}").json()["rows"][0]["review_pending"] is True
            assert client.get(f"/api/v1/tasks/{task['id']}/file").content == original
            cleared = client.get(f"/api/v1/tables/{table_id}/export.json").json()
            assert REVIEW_KEY not in cleared[0]
        assert client.delete(f"/api/v1/tasks/{task['id']}").status_code == 200
        assert client.get(f"/api/v1/tables/{table_id}").json()["rows"][0]["review_pending"] is (not resolve_review)
        assert json.loads(client.get(f"/api/v1/tables/{imported_id}/export.json").content)[0][REVIEW_KEY] is True
