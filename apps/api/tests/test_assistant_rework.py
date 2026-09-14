import json
from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from test_assistant import client as ledger_fixture, ScriptedProvider, select_profile, send, finish

from document_pipeline_api.models import (
    DataRowRecord,
    DataTableRecord,
    TaskRecord,
    AssistantToolCall,
)
from document_pipeline_api.services.assistant_tools import (
    Context,
    execute_tool,
    scoped_tool_specs,
    permit_read,
)
from document_pipeline_api.services.assistant_memory import (
    compact_history,
    summarize_result,
    encode,
)


client = ledger_fixture


def plan_run(client, operations, **context):
    ScriptedProvider.requests = []
    ScriptedProvider.plan = [
        [
            {
                "type": "tool",
                "id": "plan",
                "name": "propose_operations",
                "arguments": {"title": "本次修改", "operations": operations},
            }
        ]
    ]
    client.app.state.conversation_factory = ScriptedProvider
    return finish(
        client,
        send(
            client,
            select_profile(client),
            text="请修改数据",
            context={"table_ids": ["ledger"], **context},
        ),
    )


def decide(client, identifier):
    return client.post(f"/api/v1/assistant/tools/{identifier}/decision", json={"approve": True})


def test_batch_shared_fields_undo_and_no_duplicate_application(client):
    detail = plan_run(
        client,
        [
            {
                "kind": "update_rows",
                "table_id": "ledger",
                "row_ids": [1],
                "changes": {"total": "100"},
            },
            {
                "kind": "update_rows",
                "table_id": "ledger",
                "row_ids": [3],
                "changes": {"amount": 80},
            },
        ],
    )
    tool = detail["tools"][0]
    assert tool["status"] == "pending", detail
    assert len(tool["result"]["items"][0]["affected"]) == 2
    outcome = decide(client, tool["id"])
    assert outcome.status_code == 200, outcome.text
    assert outcome.json()["result"]["execution"]["can_undo"]
    assert decide(client, tool["id"]).json() == outcome.json()
    preview = client.post(f"/api/v1/assistant/tools/{tool['id']}/undo")
    assert preview.status_code == 200, preview.text
    assert client.post(f"/api/v1/assistant/tools/{tool['id']}/undo").json() == preview.json()
    undone = decide(client, preview.json()["id"])
    assert undone.status_code == 200, undone.text
    with client.app.state.session_factory() as session:
        assert json.loads(session.get(DataRowRecord, 1).row_json)["total"] == "1,200.50"
        assert json.loads(session.get(DataRowRecord, 2).row_json)["total"] == "1,200.50"
        assert json.loads(session.get(DataRowRecord, 3).row_json)["amount"] == 90
        assert session.get(AssistantToolCall, tool["id"]).status == "undone"


def test_new_related_row_invalidates_whole_plan(client):
    detail = plan_run(
        client,
        [
            {"kind": "update_rows", "table_id": "ledger", "row_ids": [3], "changes": {"amount": 1}},
            {
                "kind": "update_rows",
                "table_id": "ledger",
                "row_ids": [1],
                "changes": {"total": "100"},
            },
        ],
    )
    with client.app.state.session_factory() as session:
        session.add(
            DataRowRecord(
                table_id="ledger",
                item_index=9,
                row_json=json.dumps({"__row_group": "first", "total": "1,200.50"}),
            )
        )
        session.commit()
    response = decide(client, detail["tools"][0]["id"])
    assert response.status_code == 409, response.text
    with client.app.state.session_factory() as session:
        assert json.loads(session.get(DataRowRecord, 3).row_json)["amount"] == 90


def test_undo_refuses_later_user_edit(client):
    detail = plan_run(
        client,
        [{"kind": "update_rows", "table_id": "ledger", "row_ids": [3], "changes": {"amount": 1}}],
    )
    tool = detail["tools"][0]
    assert decide(client, tool["id"]).status_code == 200
    with client.app.state.session_factory() as session:
        session.get(DataRowRecord, 3).row_version += 1
        session.commit()
    assert client.post(f"/api/v1/assistant/tools/{tool['id']}/undo").status_code == 409


@pytest.mark.parametrize("context", [{"mode": "read"}, {"capabilities": ["templates"]}])
def test_permissions_enforced_even_if_model_invents_tool(client, context):
    detail = plan_run(
        client,
        [{"kind": "update_rows", "table_id": "ledger", "row_ids": [3], "changes": {"amount": 1}}],
        **context,
    )
    assert detail["tools"][0]["status"] == "failed", detail
    with client.app.state.session_factory() as session:
        assert json.loads(session.get(DataRowRecord, 3).row_json)["amount"] == 90


def test_delegation_executes_reversible_edits_but_deletion_requires_confirmation(client):
    detail = plan_run(
        client,
        [{"kind": "update_rows", "table_id": "ledger", "row_ids": [3], "changes": {"amount": 1}}],
        mode="delegate",
    )
    assert detail["tools"][0]["status"] == "approved", detail
    assert len(ScriptedProvider.requests) == 1
    assert detail["tools"][0]["result"]["execution"]["can_undo"]
    detail = plan_run(
        client, [{"kind": "delete_rows", "table_id": "ledger", "row_ids": [3]}], mode="delegate"
    )
    assert detail["tools"][0]["status"] == "pending", detail
    with client.app.state.session_factory() as session:
        assert session.get(DataRowRecord, 3) is not None


def test_multiple_scopes_catalog_and_read_only_advertisement(client):
    with client.app.state.session_factory() as session:
        session.add(
            DataTableRecord(
                id="other",
                name="不可访问",
                template_key="manual:x",
                template_version="1",
                document_kind="custom",
                columns_json="[]",
            )
        )
        session.commit()
        context = Context(table_ids=["ledger"], mode="read")
        assert [
            i["id"]
            for i in execute_tool(session, context, "catalog", {"kind": "tables"}, {})["items"]
        ] == ["ledger"]
        names = [
            s["function"]["name"] for s in scoped_tool_specs(session, context, {}, True, "修改模板")
        ]
        assert "analyze_data_table" in names and not any(
            "propose" in n or "draft" in n for n in names
        )
        with pytest.raises(HTTPException):
            execute_tool(session, context, "read_data_rows", {"table_id": "other"}, {})


def test_snapshot_paging_memory_and_single_total_chart(client):
    with client.app.state.session_factory() as session:
        context, artifacts = Context(table_ids=["ledger"]), {}
        rows = execute_tool(session, context, "read_data_rows", {"table_id": "ledger"}, artifacts)
        later = execute_tool(
            session,
            context,
            "read_analysis_snapshot",
            {"analysis_id": rows["analysis_id"], "offset": 2, "limit": 1},
            artifacts,
        )
        assert later["items"][0]["row_id"] == rows["rows"][2]["row_id"]
        assert summarize_result(rows)["data"]
        total = execute_tool(
            session, context, "analyze_data_table", {"table_id": "ledger"}, artifacts
        )
        chart = execute_tool(
            session,
            context,
            "render_chart",
            {
                "analysis_id": total["analysis_id"],
                "type": "line",
                "title": "趋势",
                "series": ["count:*"],
            },
            artifacts,
        )
        assert chart["chart"]["type"] == "metric" and "趋势" not in chart["chart"]["title"]
    history = [
        {"role": "user" if i % 2 == 0 else "assistant", "content": str(i) + "旧消息" * 3000}
        for i in range(20)
    ]
    compacted, changed = compact_history(history, 6500)
    assert changed and len(encode(compacted)) <= 6500
    assert "19" in compacted[-1]["content"]


@pytest.mark.parametrize("kind", ["create_template", "update_template", "restore_template"])
def test_chat_template_writes_removed_but_native_template_remains(client, kind):
    body = {
        "name": "合成模板",
        "fields": [{"key": "amount", "label": "金额", "value_type": "number"}],
    }
    created = client.post("/api/v1/templates", json=body)
    assert created.status_code == 201, created.text
    template = created.json()
    detail = plan_run(client, [{"kind": kind, "template_id": template["id"],
        "template": {**body, "name": "新版名称"}}], template_ids=[template["id"]])
    tool = detail["tools"][0]
    assert tool["status"] == "failed", detail
    assert "模板页面" in tool["result"]["error"]
    assert client.get(f"/api/v1/templates/{template['id']}").json()["version"] == 1
    with client.app.state.session_factory() as session:
        specs = scoped_tool_specs(session, Context(workspace=True, capabilities=["templates"]), {}, True, "修改模板")
        assert not any(s["function"]["name"].startswith("propose_") or s["function"]["name"] == "draft_template" for s in specs)
        result = execute_tool(session, Context(template_ids=[template["id"]]), "read_resource", {"kind": "template", "id": template["id"]}, {})
        assert result


def test_export_is_saved_scope_and_excel_formulas_are_literal(client):
    with client.app.state.session_factory() as session:
        row = session.get(DataRowRecord, 1)
        row.row_json = json.dumps({"vendor": "=1+1", "total": 1})
        session.commit()
    ScriptedProvider.plan = [
        [
            {
                "type": "tool",
                "id": "rows",
                "name": "read_data_rows",
                "arguments": {"table_id": "ledger"},
            }
        ]
    ]
    client.app.state.conversation_factory = ScriptedProvider
    detail = finish(
        client, send(client, select_profile(client), context={"table_id": "ledger", "row_ids": [1]})
    )
    identifier = detail["tools"][0]["result"]["analysis_id"]
    ScriptedProvider.plan = [
        [
            {
                "type": "tool",
                "id": "export",
                "name": "prepare_export",
                "arguments": {"analysis_id": identifier},
            }
        ]
    ]
    detail = finish(
        client,
        send(
            client,
            select_profile(client),
            thread_id=detail["id"],
            context={"table_id": "ledger", "row_ids": [1]},
        ),
    )
    tool = next(t for t in detail["tools"] if t["name"] == "prepare_export")
    result = client.get(f"/api/v1/assistant/tools/{tool['id']}/export")
    assert result.status_code == 200, result.text[:100] if result.status_code != 200 else ""
    workbook = load_workbook(BytesIO(result.content))
    assert workbook.active.max_row == 2
    assert workbook.active["A2"].value == "=1+1" and workbook.active["A2"].data_type == "s"


def test_authorized_source_page_and_pending_file_row_count(client, tmp_path):
    from document_pipeline_api.services.assistant_analysis import analyze, AnalysisRequest

    settings = client.app.state.settings
    settings.storage_dir.mkdir(exist_ok=True, parents=True)
    path = settings.storage_dir / "source.txt"
    path.write_text("合成原件 " * 1200, encoding="utf-8")
    with client.app.state.session_factory() as session:
        task = TaskRecord(
            id="source",
            sha256="a" * 64,
            filename="合成原件.txt",
            content_type="text/plain",
            storage_path="source.txt",
            size_bytes=path.stat().st_size,
            status="needs_review",
            template_mode="auto",
        )
        session.add(task)
        session.flush()
        for identifier in [1, 2, 3]:
            row = session.get(DataRowRecord, identifier)
            row.task_id = task.id
            row.review_pending = True
        session.get(DataRowRecord, 4).review_pending = False
        session.commit()
        context = Context(table_id="ledger", row_ids=[1])
        permit_read(session, context, "task", "source")
        original = execute_tool(
            session, context, "read_original_page", {"task_id": "source"}, {}, settings
        )
        assert original["text"].startswith("合成原件") and original["next_offset"] == 4000
        with pytest.raises(HTTPException):
            permit_read(session, context, "task", "not-selected")
        result = analyze(session, AnalysisRequest(table_id="ledger"))
        assert result["source"]["review_pending_documents"] == 1
        assert result["source"]["review_pending_count"] == 3
        assert any("1 份来源" in w and "3 条明细" in w for w in result["warnings"])


def test_task_batch_and_table_operations_use_native_services(client):
    with client.app.state.session_factory() as session:
        for identifier in ["task-a", "task-b"]:
            session.add(
                TaskRecord(
                    id=identifier,
                    sha256=identifier * 8,
                    filename=identifier + ".txt",
                    content_type="text/plain",
                    storage_path=identifier + ".txt",
                    size_bytes=10,
                    status="queued",
                    template_mode="auto",
                    model_name="original-model",
                    model_provider="lm_studio",
                )
            )
        session.commit()
    detail = plan_run(
        client,
        [{"kind": "tasks", "task_ids": ["task-a", "task-b"], "action": "pause"}],
        task_ids=["task-a", "task-b"],
    )
    tool = detail["tools"][0]
    assert tool["status"] == "pending", detail
    response = decide(client, tool["id"])
    assert response.status_code == 200, response.text
    with client.app.state.session_factory() as session:
        for identifier in ["task-a", "task-b"]:
            task = session.get(TaskRecord, identifier)
            assert task.status == "paused" and task.model_name == "original-model"
    detail = plan_run(
        client,
        [
            {"kind": "rename_table", "table_id": "ledger", "name": "重新命名"},
            {"kind": "add_column", "table_id": "ledger", "name": "备注"},
            {"kind": "split_table", "table_id": "ledger", "field_key": "vendor"},
        ],
    )
    response = decide(client, detail["tools"][0]["id"])
    assert response.status_code == 200, response.text
    assert len(response.json()["result"]["execution"]["results"][-1]["views"]) == 3


def test_native_failure_rolls_back_earlier_operations(client, monkeypatch):
    detail = plan_run(
        client,
        [
            {"kind": "update_rows", "table_id": "ledger", "row_ids": [3], "changes": {"amount": 1}},
            {"kind": "rename_table", "table_id": "ledger", "name": "后来失败"},
        ],
    )

    def fail(*args, **kwargs):
        raise HTTPException(409, "模拟服务冲突")

    monkeypatch.setattr("document_pipeline_api.services.data_tables.rename_table", fail)
    assert decide(client, detail["tools"][0]["id"]).status_code == 409
    with client.app.state.session_factory() as session:
        assert json.loads(session.get(DataRowRecord, 3).row_json)["amount"] == 90
        assert session.get(AssistantToolCall, detail["tools"][0]["id"]).status == "stale"


def test_original_image_is_given_to_model_but_not_saved_as_base64(client):
    from PIL import Image

    settings = client.app.state.settings
    settings.storage_dir.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (100, 100), "white").save(settings.storage_dir / "image.png")
    with client.app.state.session_factory() as session:
        session.add(
            TaskRecord(
                id="image",
                sha256="b" * 64,
                filename="image.png",
                content_type="image/png",
                storage_path="image.png",
                size_bytes=100,
                status="completed",
                template_mode="auto",
            )
        )
        session.commit()
    ScriptedProvider.requests = []
    ScriptedProvider.plan = [
        [
            {
                "type": "tool",
                "id": "original",
                "name": "read_original_page",
                "arguments": {"task_id": "image", "vision": True},
            }
        ]
    ]
    client.app.state.conversation_factory = ScriptedProvider
    detail = finish(client, send(client, select_profile(client), context={"task_ids": ["image"]}))
    assert detail["tools"][0]["status"] == "completed", detail
    assert any("data:image/jpeg;base64," in json.dumps(m) for m in ScriptedProvider.requests[-1])
    assert "base64" not in json.dumps(detail)


def test_local_enum_null_compatibility_and_on_demand_operation_schema(client):
    with client.app.state.session_factory() as session:
        context, artifacts = Context(table_ids=["ledger"]), {}
        result = execute_tool(
            session,
            context,
            "analyze_data_table",
            {"table_id": "ledger", "time_bucket": "null"},
            artifacts,
        )
        assert result["source"]["request"]["time_bucket"] is None
        names = [
            s["function"]["name"]
            for s in scoped_tool_specs(session, context, {}, False, "把这笔设成一百")
        ]
        assert "get_operation_tools" in names
        execute_tool(session, context, "get_operation_tools", {"category": "data"}, artifacts)
        names = [
            s["function"]["name"]
            for s in scoped_tool_specs(session, context, artifacts, False, "把这笔设成一百")
        ]
        assert "propose_operations" in names


def test_full_tool_results_cannot_be_recalled_after_scope_is_removed(client):
    client.app.state.conversation_factory = ScriptedProvider
    ScriptedProvider.plan = [
        [
            {
                "type": "tool",
                "id": "read",
                "name": "read_data_rows",
                "arguments": {"table_id": "ledger"},
            }
        ]
    ]
    detail = finish(client, send(client, select_profile(client), context={"table_ids": ["ledger"]}))
    previous = detail["tools"][0]["id"]
    ScriptedProvider.plan = [
        [
            {
                "type": "tool",
                "id": "recall",
                "name": "read_tool_result",
                "arguments": {"tool_id": previous},
            }
        ]
    ]
    detail = finish(
        client, send(client, select_profile(client), thread_id=detail["id"], context={})
    )
    recall = next(t for t in detail["tools"] if t["name"] == "read_tool_result")
    assert recall["status"] == "failed" and "授权范围" in recall["result"]["error"]


@pytest.mark.parametrize("ollama", [True, False])
def test_native_image_protocol_conversion(ollama, tmp_path):
    import threading
    import httpx
    from document_pipeline_api.config import Settings
    from document_pipeline_api.model_providers.conversation import ConversationProvider

    def handler(request):
        payload = json.loads(request.content)
        message = payload["messages"][0]
        if ollama:
            assert message["content"] == "合成图片" and message["images"] == ["aGVsbG8="]
            return httpx.Response(
                200, text=json.dumps({"message": {"content": "已收到"}, "done": True}) + "\n"
            )
        assert message["content"][1]["image_url"]["url"] == "data:image/png;base64,aGVsbG8="
        return httpx.Response(
            200,
            text="data: "
            + json.dumps({"choices": [{"delta": {"content": "已收到"}}]})
            + "\n\ndata: [DONE]\n\n",
        )

    provider = ConversationProvider(
        Settings(
            database_url="sqlite:///:memory:",
            storage_dir=tmp_path,
            model_provider="ollama" if ollama else "lm_studio",
            model_base_url="http://127.0.0.1:1234",
            model_name="synthetic",
        ),
        threading.Event(),
    )
    provider.client.close()
    provider.client = httpx.Client(transport=httpx.MockTransport(handler))
    try:
        result = list(
            provider.stream(
                [
                    {
                        "role": "user",
                        "content": [
                            {"type": "text", "text": "合成图片"},
                            {
                                "type": "image_url",
                                "image_url": {"url": "data:image/png;base64,aGVsbG8="},
                            },
                        ],
                    }
                ],
                [],
            )
        )
        assert any(e.get("text") == "已收到" for e in result)
    finally:
        provider.close()
